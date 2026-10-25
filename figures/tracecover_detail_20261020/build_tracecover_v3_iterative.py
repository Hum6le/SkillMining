"""TraceCover v3: feasible initialization plus the implemented iterative search.

All artwork is editable native PowerPoint vector/text. Graphs are schematic;
their purpose is to show the state and control flow of the actual solver.
"""
from __future__ import annotations

import ast
import json
import math
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'.pptx_deps'))
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

C=dict(ink='182235',black='111111',blue='42B9ED',navy='4678C4',pale_blue='E8F3FC',
       cream='FFF6CE',yellow='FFD75F',coral='F36F63',teal='59CCAF',pale_teal='E4F7EC',
       purple='ECE3F7',gray='E3E5E8',paper='F7F8FA',muted='667385',green='339866',
       red='D65555',light='D2D7DD',white='FFFFFF',neon='F6EE4B',mint='8DE5A6',
       pink='F8DDE2',steel='C4CCD6')
def load_defs(filename,names):
    tree=ast.parse((OUT/filename).read_text(encoding='utf-8'))
    items=[n for n in tree.body if isinstance(n,(ast.FunctionDef,ast.ClassDef)) and n.name in names]
    exec(compile(ast.Module(body=items,type_ignores=[]),filename,'exec'),globals())
load_defs('build_tracecover_figure.py',{'Canvas','rgb'})
load_defs('build_tracecover_v2.py',{'marker','note','polygon','highlighter','scissors','zoom','swap_icon'})

W,H=21.1,10.65
PPTX=OUT/'TraceCover_iterative_editable_v3.pptx'
prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
slide=prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb('white')
c=Canvas(slide.shapes)


def heading(text,x,y,w,number):
    c.circle(x,y,.40,'cream','ink',1)
    c.text(str(number),x,y+.03,.40,.32,16,True,PP_ALIGN.CENTER)
    marker(c,[(x+.58,y+.32),(x+w-.10,y+.28)],'neon',15,60000)
    c.text(text,x+.53,y-.015,w-.53,.46,21,True)


def flow(points,color='coral',width=2.6,dash=False):
    return c.path(points,color,width,True,dash)


def diamond(x,y,w,h,lines,fill='cream',size=14):
    s=c.shape(MSO_SHAPE.DIAMOND,x,y,w,h,fill,'ink',1.2)
    c.text(lines,x+w*.18,y+h*.17,w*.64,h*.66,size,True,PP_ALIGN.CENTER)
    return s


def clock(x,y,d=.38):
    c.circle(x,y,d,'white','ink',1.2)
    c.line(x+d*.5,y+d*.5,x+d*.5,y+d*.19,'ink',1.2)
    c.line(x+d*.5,y+d*.5,x+d*.75,y+d*.62,'ink',1.2)


def cross(x,y,s=.22):
    c.circle(x,y,s,'pink','red',.7)
    c.line(x+s*.25,y+s*.25,x+s*.75,y+s*.75,'red',1.1)
    c.line(x+s*.25,y+s*.75,x+s*.75,y+s*.25,'red',1.1)


POS={'a':(.17,.91),'b':(1.12,.91),'c':(2.12,.18),'d':(3.13,.18),
     'e':(2.12,1.64),'f':(3.13,1.64)}
ICONS={'a':'account','b':'validate','c':'status','d':'refund','e':'reason','f':'update'}
OLD_TREE=[('a','b'),('b','c'),('c','d'),('b','e'),('e','f')]
OLD_RES=[('c','f'),('d','f'),('e','e')]
NEW_TREE=[('a','b'),('b','c'),('c','d'),('b','e'),('c','f')]
NEW_RES=[('d','f'),('e','e')]


def graph(ox,oy,scale=1.,mode='current',show_removed=False):
    pts={k:(ox+xx*scale,oy+yy*scale) for k,(xx,yy) in POS.items()}
    rad=.23*scale
    tree=NEW_TREE if mode in {'updated','output','recompute'} else OLD_TREE
    residual=NEW_RES if mode in {'updated','output','recompute'} else ([] if mode=='seed' else OLD_RES)
    def draw_edge(u,v,color,width,dash=False,highlight=False):
        xx,yy=pts[u];tx,ty=pts[v]
        if u==v:
            pp=[(xx-.14*scale,yy+.18*scale),(xx-.34*scale,yy+.40*scale),
                (xx-.34*scale,yy+.72*scale),(xx+.34*scale,yy+.72*scale),
                (xx+.34*scale,yy+.40*scale),(xx+.14*scale,yy+.18*scale)]
            if mode=='output':marker(c,pp,'mint',10,45000)
            c.path(pp,color,width,True,dash)
        else:
            dx,dy=tx-xx,ty-yy;dist=math.hypot(dx,dy)
            aa=(xx+rad*dx/dist,yy+rad*dy/dist);bb=(tx-rad*dx/dist,ty-rad*dy/dist)
            if highlight:marker(c,[aa,bb],'neon',13*scale,65000)
            c.line(*aa,*bb,color,width,True,dash)
    if show_removed:draw_edge('e','f','light',1.1,True)
    for u,v in residual:draw_edge(u,v,'navy',1.6 if scale>.8 else 1.2,True)
    for u,v in tree:draw_edge(u,v,'ink',1.65 if scale>.8 else 1.15,False,True)
    ax,ay=pts['a']
    c.circle(ax-.048*scale,ay-.57*scale,.096*scale,'paper','muted',.6)
    c.line(ax,ay-.47*scale,ax,ay-rad,'muted',.9,True,True)
    for k,(xx,yy) in pts.items():
        c.circle(xx-rad,yy-rad,rad*2,'cream' if mode in {'updated','output'} else ('purple' if k=='f' else 'pale_blue'),'ink',1)
        ss=.30*scale;c.icon(ICONS[k],xx-.46*ss,yy-.42*ss,ss,'cream' if k=='e' else 'pale_blue')
        label_y=yy-rad-.26*scale if mode=='current' and k=='b' else yy+rad+.015*scale
        c.text(k,xx-.24*scale,label_y,.48*scale,.21*scale,max(10.5,12*scale),True,PP_ALIGN.CENTER)
    return pts


c.text('TraceCover',.36,.16,3.70,.48,30,True)
c.text('Feasible initialization + iterative graph–backbone refinement',3.89,.19,16.72,.41,24)
heading('Feasible initialization',.40,.91,4.08,1)
heading('Iterative joint refinement',5.15,.91,10.92,2)
heading('Selected structure',17.38,.91,3.28,3)
# Soft region cues; the center remains open to make the control loop prominent.
c.line(4.84,1.66,4.84,9.91,'light',1,False,True)
c.line(17.03,1.72,17.03,9.91,'light',1,False,True)

# 1. Initialization starts from an evidence-weighted maximum rooted tree.
for j in [2,1,0]:
    p=c.rect(.79+j*.08,1.69-j*.055,3.27,.62,'paper' if j else 'cream','ink',.8)
    p.rotation=(j-1)*1.2
c.icon('chat',.95,1.86,.35,'pale_blue')
c.text('Observed graph + trace evidence',1.43,1.80,2.42,.35,13.1,True)
c.arrow(2.22,2.40,.25,.25,'down')
c.text('Maximum-evidence seed B₀',.83,2.78,3.18,.29,16,True,PP_ALIGN.CENTER)
graph(.87,3.28,.86,'seed')
flow([(2.45,5.01),(2.45,5.30)],'coral',2)
diamond(1.52,5.34,1.84,1.06,'Fρ(G) ≥ η?',size=14)
c.text('Yes',3.41,5.53,.48,.25,13,True)
flow([(3.36,5.87),(4.58,5.87),(4.58,3.57),(5.27,3.57)],'coral',2.6)
c.math(['G',('0',-24000)],3.84,5.42,.50,.30,19)
c.text('No',2.63,6.47,.43,.23,12,True)
flow([(2.45,6.40),(2.45,6.83)],'coral',2)
c.rect(.98,6.91,3.15,.91,'pale_blue',None)
c.circle(1.18,7.13,.32,'mint','ink',.9)
c.text('+',1.18,7.12,.32,.29,20,True,PP_ALIGN.CENTER)
c.text('Add best residual edge',1.67,6.99,2.19,.28,14.2,True)
c.math(['maximize  ΔF',('ρ',-24000),' / ΔC'],1.64,7.37,2.23,.25,15)
flow([(1.06,7.36),(.51,7.36),(.51,5.87),(1.49,5.87)],'navy',2.6)
note(c,'add → update coverage',.88,8.00,3.38,.35,13,angle=-3)
c.text('Repeat until coverage is sufficient.',.89,8.46,3.39,.38,12.1,align=PP_ALIGN.CENTER)
cross(1.07,9.04,.25)
c.text('No useful addition and Fρ < η:\nmark initialization infeasible.',1.52,8.95,2.68,.55,11.9)

# 2. Current state. The prominent accepted-update arrow returns here.
marker(c,[(5.46,2.09),(8.72,2.04)],'neon',13,49000)
c.math(['Current  G',('k',-24000),', B',('k',-24000)],5.38,1.80,3.45,.35,22)
graph(5.49,2.81,.98,'current')
c.text('Current graph + current backbone',5.17,2.32,3.70,.26,12,align=PP_ALIGN.CENTER)

# Check budget before forming and evaluating this iteration's proposals.
diamond(9.10,3.08,1.47,.98,'Budget\nleft?',fill='pale_blue',size=12)
clock(9.64,2.60,.34)
flow([(8.88,3.57),(9.07,3.57)],'coral',2.3)
c.text('Yes',10.60,3.19,.40,.23,12,True)
flow([(10.58,3.57),(10.91,3.57)],'coral',2.3)
flow([(9.84,3.05),(9.84,1.52),(16.73,1.52),(16.73,2.35),(17.42,2.35)],'muted',1.8)
c.text('No: move budget exhausted',11.08,1.20,4.53,.25,12,align=PP_ALIGN.CENTER)

# Candidate edit cards. The graph data and eligible observed-edge pool stay fixed.
c.text('Propose local graph edits',10.93,1.87,5.34,.31,18,True,PP_ALIGN.CENTER)
for xx,fill,angle in [(10.98,'cream',-1.5),(13.69,'pale_blue',1.5)]:
    p=c.rect(xx,2.40,2.48,1.58,fill,None);p.rotation=angle
scissors(c,11.19,2.64,.56)
c.text('Delete',11.94,2.55,1.33,.29,16,True)
c.math(['E′ = E ∖ {e}'],11.18,3.51,2.07,.27,16)
# A crossed-out example edge, under the delete pictogram.
c.line(12.02,3.07,12.90,3.07,'light',1.7,True)
c.line(12.32,2.95,12.57,3.20,'red',1.5)
c.line(12.32,3.20,12.57,2.95,'red',1.5)
swap_icon(c,13.98,2.70,.71)
c.text('Exchange',14.66,2.55,1.32,.29,16,True)
c.math(['E′ = (E ∖ {e}) ∪ {e⁺}'],13.77,3.51,2.31,.27,13.8)
c.text('unused observed edge',14.72,3.03,1.27,.31,10.1,align=PP_ALIGN.CENTER)
flow([(13.56,4.05),(13.56,4.32)],'coral',2.3)

# Fast candidate screening: objective lower bound plus coverage.
c.text('Screen bounds & trace coverage',11.66,4.35,4.80,.31,17,True,PP_ALIGN.CENTER)
for i,(xx,yy) in enumerate([(12.28,4.86),(12.82,4.73),(13.34,4.88)]):
    p=c.rect(xx,yy,.44,.42,'paper','ink',.6);p.rotation=(-5,3,-3)[i]
    for z in range(2):c.line(xx+.09,yy+.13+z*.12,xx+.35,yy+.13+z*.12,'muted',.7)
polygon(c,[(12.15,5.32),(13.94,5.32),(13.29,5.77),(13.29,6.00),
           (12.81,6.00),(12.81,5.77)],'pale_blue','ink',1.0)
c.text('bounded shortlist',14.14,4.88,2.01,.25,12.5,True)
c.check(14.17,5.31,.25)
c.math(['F',('ρ',-24000),'(G′) ≥ η'],14.55,5.27,1.55,.29,14.3)
cross(14.17,5.74,.25)
c.text('bound cannot improve J',14.56,5.72,1.73,.36,10.4)
flow([(13.06,6.04),(13.06,6.35)],'coral',2.3)

# A small objective note inside the loop replaces the former dominant formula box.
p=c.rect(8.65,4.79,2.60,1.50,'cream',None);p.rotation=-2
marker(c,[(8.88,5.69),(10.99,5.66)],'neon',11,47000)
c.math(['J(G) = C(G) − βW',('tree',-24000),'(G)'],8.80,5.03,2.29,.41,16.5)
c.text('Compare candidates\nunder fixed training evidence',8.85,5.69,2.20,.44,11.2,align=PP_ALIGN.CENTER)

# Exact evaluation is bounded to the shortlisted graphs.
c.text('Recompute backbone & score shortlist',11.78,6.40,4.69,.30,17,True,PP_ALIGN.CENTER)
highlighter(c,11.85,6.92,.56)
graph(12.94,6.89,.62,'recompute',show_removed=True)
c.math(['B′ = MaxTree(G′)'],11.74,8.20,2.69,.25,15.5)
c.text('Check root reachability; compare J(G′).',12.17,8.58,4.10,.25,11.8,align=PP_ALIGN.CENTER)
# Arrow exits left from the evaluation region and enters the decision gate.
flow([(12.12,7.65),(10.45,7.65),(10.45,8.18)],'coral',2.3)
c.text('Best feasible\nshortlisted move',10.19,6.86,1.55,.45,10.8,align=PP_ALIGN.CENTER)

# Only a round with an improving candidate leads to another iteration.
diamond(9.25,8.24,2.39,1.28,'Improvement\nfound?',size=14)
c.text('Yes',8.88,8.42,.37,.25,12,True)
flow([(9.23,8.88),(8.82,8.88)],'coral',2.6)
# If all shortlisted candidates fail to improve, return the current graph.
flow([(11.67,8.88),(16.73,8.88),(16.73,2.35),(17.42,2.35)],'muted',1.8)
c.text('No: no feasible improving candidate',12.01,9.06,4.47,.30,12.1,align=PP_ALIGN.CENTER)

# Update card illustrates that the backbone may be revised, too.
c.rect(5.22,7.15,3.64,2.52,'pale_teal',None)
c.check(5.41,7.34,.29)
c.math(['Accept best → G',('k+1',-24000),', B',('k+1',-24000)],5.85,7.26,2.79,.35,17)
graph(5.56,7.88,.80,'updated',show_removed=True)
flow([(6.78,7.11),(6.78,4.57),(6.59,3.97)],'navy',4)
note(c,'next\niteration',5.20,5.64,1.44,.72,14.5,angle=-3)
c.text('Re-highlight the backbone after each update.',5.23,9.88,3.67,.29,12.1,align=PP_ALIGN.CENTER)

# Accepted-move guarantees are shown at the bottom of the main cycle.
marker(c,[(9.43,10.00),(16.48,9.96)],'mint',21,40000)
c.math(['Accepted:  F',('ρ',-24000),'(G',('k+1',-24000),') ≥ η,   J(G',('k+1',-24000),') < J(G',('k',-24000),')'],9.37,9.79,7.15,.37,17)

# 3. Both stop exits return the current selected graph and backbone.
c.shape(MSO_SHAPE.OCTAGON,17.49,2.09,.51,.51,'pink','ink',1.0)
c.rect(17.63,2.32,.23,.055,'red',None,rounded=False)
c.text('Stop & return',18.13,2.15,2.29,.36,18,True)
c.arrow(18.80,2.83,.26,.28,'down')
graph(17.61,3.59,.87,'output')
c.math(['G, B, E',('R',-24000)],17.66,6.26,2.98,.40,24)
note(c,'spine + branches + retries',17.29,6.85,3.54,.37,13.5,angle=-3)
c.line(18.96,7.48,18.96,7.91,'muted',1.6,True)
c.icon('doc',18.58,8.09,.74,'cream')
c.text('Evidence induction\n& text-skill compilation',17.52,9.04,3.15,.61,16,True,PP_ALIGN.CENTER)
c.text('Downstream of TraceCover',17.57,9.79,3.06,.25,11.7,align=PP_ALIGN.CENTER)

# Compact visual legend and scope note; no separate legend panel.
c.line(.56,10.27,.88,10.27,'navy',1.7,True,True)
c.text('retained residual',.97,10.13,1.69,.27,11.5)
marker(c,[(2.93,10.27),(3.25,10.27)],'neon',10,65000)
c.line(2.93,10.27,3.25,10.27,'ink',1.1,True)
c.text('backbone',3.34,10.13,1.02,.27,11.5)
c.text('Schematic graph states · fixed observed edge pool · budgeted search, without an optimality guarantee',5.31,10.27,15.19,.23,11.1,align=PP_ALIGN.RIGHT)

slide.notes_slide.notes_text_frame.text=(
    'TraceCover actual-solver illustration, iterative v3. All shapes, markers, pictograms, '
    'arrows and equations are native editable PowerPoint objects. '
    'The initialization first computes a maximum evidence-weighted rooted arborescence. '
    'It repeatedly adds a support-qualified observed residual with maximal marginal '
    'saturated-coverage gain per marginal complexity cost until the corpus target is met. '
    'If that target remains unattainable, it marks the initialization infeasible. '
    'The main loop retains a current graph and a current maximum rooted spanning backbone. '
    'A move budget is checked before proposing and evaluating another round. '
    'Candidates delete one selected edge or exchange it for one eligible unused observed edge. '
    'Bound and coverage screening yields a bounded shortlist. Each required maximum '
    'arborescence is recomputed in the candidate graph; root-unreachable candidates fail. '
    'Deleting a non-tree edge alone can reuse the current backbone, as in the implementation. '
    'The best feasible strictly improving candidate is accepted each round. '
    'No shortlisted improvement, an empty screened neighborhood, or exhausted move budget '
    'returns the current state. This is not a test for global or unrestricted local optimality. '
    'The bottom accepted-state example removes e->f, promoting the already retained c->f '
    'to the backbone; d->f and e->e remain residual. This is a schematic possible move, '
    'not a run result or a complete ABCD SOP. All six action nodes remain root-reachable. '
    'Only the joint objective must decrease along accepted moves; C, W and F need not '
    'each improve monotonically. Selected graph and raw evidence go to later induction '
    'and compilation; the fixed agent eventually consumes a text skill, not this graph.')

prs.save(PPTX)
import networkx as nx
assert nx.is_arborescence(nx.DiGraph(OLD_TREE))
assert nx.is_arborescence(nx.DiGraph(NEW_TREE))
assert set(NEW_TREE+NEW_RES)==set(OLD_TREE+OLD_RES)-{('e','f')}
assert set(nx.descendants(nx.DiGraph(NEW_TREE),'a'))==set(POS)-{'a'}
check=Presentation(PPTX)
def count_shapes(shapes):
    n=0
    for sh in shapes:
        assert sh.shape_type!=13
        n+=count_shapes(sh.shapes) if sh.shape_type==6 else 1
    return n
count=count_shapes(check.slides[0].shapes)
for sh in check.slides[0].shapes:
    assert sh.left>=0 and sh.top>=0,sh.name
    assert sh.left+sh.width<=check.slide_width+100,(sh.name,'right')
    assert sh.top+sh.height<=check.slide_height+100,(sh.name,'bottom')
manifest={'version':3,'title':'TraceCover iterative implementation','canvas_inches':[W,H],
          'editable_objects':count,'raster_objects':0,
          'sources':['skill_mining/joint_trace_cover.py','skill_mining/backbone_workflow_mining.py'],
          'initialization_loop':'max tree; repeat add best coverage gain / complexity; check feasibility',
          'main_loop':['check budget','propose deletions/exchanges','screen bounds and coverage',
                       'evaluate bounded shortlist and candidate backbone','accept best feasible improvement',
                       'update graph/backbone; repeat'],
          'stop_conditions':['move budget exhausted','no feasible strictly improving shortlisted candidate'],
          'scope':'schematic state example, not an empirical mining result',
          'illustrative_before':{'tree':OLD_TREE,'residual':OLD_RES},
          'illustrative_after':{'tree':NEW_TREE,'residual':NEW_RES}}
(OUT/'figure_manifest_v3.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(f'Created {PPTX.name}; {count} native editable objects.')


def export():
    import fitz
    import win32com.client
    app=win32com.client.DispatchEx('PowerPoint.Application');pres=None
    try:
        pres=app.Presentations.Open(str(PPTX),ReadOnly=True,Untitled=False,WithWindow=False)
        pres.SaveAs(str(OUT/'TraceCover_iterative_v3.pdf'),32)
        pres.Slides(1).Export(str(OUT/'TraceCover_iterative_v3.png'),'PNG',3798,1917)
        with fitz.open(OUT/'TraceCover_iterative_v3.pdf') as doc:
            assert len(doc)==1
            text=doc[0].get_text()
            assert 'TraceCover' in text and 'Improvement' in text and 'Budget' in text
            assert len(doc[0].get_images())==0
            print('Exported one-page vector PDF with',len(doc[0].get_drawings()),'paths and no embedded images.')
    finally:
        if pres is not None:pres.Close()
        app.Quit()


if __name__=='__main__':export()
