"""Marker-annotated TraceCover detail, with native editable PPT artwork.

Reuse the preceding figure's drawing primitives without executing its builder.
"""
from __future__ import annotations

import ast
import copy
import json
import math
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'.pptx_deps'))
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

C=dict(ink='182235',black='111111',blue='42B9ED',navy='4678C4',pale_blue='E8F3FC',
       cream='FFF6CE',yellow='FFD75F',coral='F36F63',teal='59CCAF',pale_teal='E4F7EC',
       purple='ECE3F7',gray='E3E5E8',paper='F7F8FA',muted='667385',green='339866',
       red='D65555',light='D2D7DD',white='FFFFFF',neon='F6EE4B',mint='8DE5A6',
       pink='F8DDE2',steel='C4CCD6')
source=ast.parse((OUT/'build_tracecover_figure.py').read_text(encoding='utf-8'))
defs=[n for n in source.body if isinstance(n,(ast.FunctionDef,ast.ClassDef)) and n.name in {'Canvas','rgb'}]
exec(compile(ast.Module(body=defs,type_ignores=[]),'drawing_primitives','exec'),globals())

W,H=18.6,9.5
PPTX=OUT/'TraceCover_detail_editable_v2.pptx'
prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
slide=prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb('white')
c=Canvas(slide.shapes)


def marker(cc,points,color='neon',width=16,alpha=55000):
    s=cc.path(points,color,width,False)
    ln=s.line._get_or_add_ln();ln.set('cap','rnd')
    for node in ln.xpath('./a:solidFill/a:srgbClr'):
        a=OxmlElement('a:alpha');a.set('val',str(int(alpha)));node.append(a)
    return s


def note(cc,text,x,y,w,h=.34,size=15,color='ink',angle=-3):
    s=cc.text(text,x,y,w,h,size,False,PP_ALIGN.CENTER,color,font='Segoe Print')
    s.rotation=angle
    return s


def polygon(cc,points,fill='steel',line='ink',lw=1):
    pts=[(int(Inches(cc.ox+x)),int(Inches(cc.oy+y))) for x,y in points]
    fb=cc.parent.build_freeform(*pts[0]);fb.add_line_segments(pts[1:],close=True)
    s=fb.convert_to_shape();s.fill.solid();s.fill.fore_color.rgb=rgb(fill)
    s.line.color.rgb=rgb(line);s.line.width=Pt(lw)
    s._element.spPr.append(OxmlElement('a:effectLst'))
    for ref in s._element.xpath('./p:style/a:effectRef'):ref.set('idx','0')
    return s


def highlighter(cc,x,y,s=1.):
    # Chisel nib at the bottom left; every component is editable.
    g=cc.parent.add_group_shape();gc=Canvas(g.shapes,cc.ox+x,cc.oy+y)
    gc.rect(.15*s,.03*s,.38*s,.85*s,'yellow','ink',1.2)
    gc.rect(.15*s,.09*s,.38*s,.13*s,'ink','ink',.6,False)
    gc.rect(.18*s,.34*s,.32*s,.10*s,'cream','ink',.6,False)
    gc.rect(.20*s,.87*s,.28*s,.17*s,'paper','ink',.8,False)
    polygon(gc,[(.20*s,1.04*s),(.48*s,1.04*s),(.39*s,1.25*s),(.20*s,1.17*s)],'yellow')
    g.rotation=36
    return g


def scissors(cc,x,y,s=1.):
    # Handles lower left, sharp open blades upper right.
    polygon(cc,[(x+.29*s,y+.60*s),(x+.47*s,y+.43*s),(x+1.05*s,y+.10*s),
                (x+.58*s,y+.61*s),(x+.45*s,y+.70*s)],'steel')
    polygon(cc,[(x+.39*s,y+.68*s),(x+.55*s,y+.59*s),(x+1.05*s,y+.91*s),
                (x+.55*s,y+.78*s)],'steel')
    cc.circle(x,y+.47*s,.43*s,'pale_blue','ink',1.2)
    cc.circle(x+.13*s,y+.87*s,.43*s,'pink','ink',1.2)
    cc.circle(x+.10*s,y+.57*s,.23*s,'white','ink',.8)
    cc.circle(x+.23*s,y+.97*s,.23*s,'white','ink',.8)
    cc.circle(x+.42*s,y+.61*s,.13*s,'yellow','ink',.7)


def zoom(cc,x,y,s=1.):
    cc.circle(x,y,.83*s,'white','ink',2)
    cc.line(x+.67*s,y+.66*s,x+1.16*s,y+1.14*s,'ink',5)
    pts=[(x+.19*s,y+.45*s),(x+.40*s,y+.22*s),(x+.64*s,y+.48*s)]
    marker(cc,[pts[0],pts[1],pts[2]],'neon',8,.65*100000)
    for p,q in zip(pts,pts[1:]):cc.line(*p,*q,'ink',1.1)
    for xx,yy in pts:cc.circle(xx-.06*s,yy-.06*s,.12*s,'blue','ink',.8)
    cc.text('wₑ',x+.27*s,y+.52*s,.34*s,.18*s,10,True,PP_ALIGN.CENTER)


def swap_icon(cc,x,y,s=1.):
    cc.path([(x,y+.24*s),(x+.48*s,y+.24*s),(x+.64*s,y+.40*s)],'navy',2,True)
    cc.path([(x+.65*s,y+.70*s),(x+.17*s,y+.70*s),(x,y+.54*s)],'navy',2,True)


def heading(text,x,w,num):
    c.circle(x,.88,.39,'cream','ink',1)
    c.text(str(num),x,.91,.39,.30,16,True,PP_ALIGN.CENTER)
    marker(c,[(x+.57,1.18),(x+w-.02,1.12)],'neon',16,62000)
    c.text(text,x+.53,.84,w-.53,.46,22,True)


c.text('TraceCover',.38,.16,3.60,.47,30,True)
c.text('Find a shared spine. Keep supported alternatives.',3.83,.20,12.88,.39,23)
c.text('OFFLINE',16.95,.24,1.24,.30,13,True,PP_ALIGN.RIGHT)
heading('Evidence-weighted graph',.45,5.01,1)
heading('Highlight, test & revise',6.47,5.14,2)
heading('Spine + useful deviations',12.69,5.33,3)

# Loose paper stack, with action tokens rather than a text-heavy box.
for j in [2,1,0]:
    p=c.rect(.70+j*.10,1.54-j*.065,4.40,.70,'paper' if j else 'cream','ink',.85)
    p.rotation=(j-1)*1.4
c.icon('chat',.88,1.70,.40,'pale_blue')
c.text('Trace 1',1.47,1.60,.95,.25,14,True)
for j,k in enumerate(['account','validate','status','refund']):
    xx=2.52+j*.53
    if j:c.line(xx-.16,1.85,xx-.03,1.85,'muted',.9,True)
    c.icon(k,xx,1.67,.28)
c.text('…',4.81,1.59,.37,.35,22,True)
c.arrow(2.87,2.33,.25,.27,'down')

NODES={'account':(.40,1.13),'validate':(1.50,1.13),'status':(2.80,.36),
       'refund':(4.05,.36),'reason':(2.80,1.98),'update':(4.05,1.98)}
TREE=[('account','validate'),('validate','status'),('status','refund'),
      ('validate','reason'),('reason','update')]
RES=[('status','update'),('reason','reason')]
LABEL={'account':'Account','validate':'Validate','status':'Status','refund':'Refund',
       'reason':'Reason','update':'Update'}


def graph(cc,ox,oy,mode='raw',scale=1.):
    pts={k:(ox+x*scale,oy+y*scale) for k,(x,y) in NODES.items()}
    rad=.25*scale
    def edge(u,v,color='ink',width=1.5,dash=False,highlight=False):
        x1,y1=pts[u];x2,y2=pts[v]
        if u==v:
            pp=[(x1-.17*scale,y1+.18*scale),(x1-.39*scale,y1+.42*scale),
                (x1-.39*scale,y1+.78*scale),(x1+.38*scale,y1+.78*scale),
                (x1+.38*scale,y1+.42*scale),(x1+.17*scale,y1+.18*scale)]
            if highlight:marker(cc,pp,'mint',12,42000)
            cc.path(pp,color,width,True,dash)
        else:
            dx,dy=x2-x1,y2-y1;d=math.hypot(dx,dy)
            a=(x1+rad*dx/d,y1+rad*dy/d);b=(x2-rad*dx/d,y2-rad*dy/d)
            if highlight:marker(cc,[a,b],'neon',15,65000)
            cc.line(*a,*b,color,width,True,dash)
    if mode!='final':
        for u,v,dx in [('refund','update',.30),('update','refund',.52)]:
            xx,yy=pts[u];tx,ty=pts[v];sgn=1 if yy<ty else -1
            pp=[(xx+.21*scale,yy+.12*scale*sgn),
                (xx+dx*scale,yy+.12*scale*sgn),
                (tx+dx*scale,ty-.12*scale*sgn),(tx+.21*scale,ty-.12*scale*sgn)]
            cc.path(pp,'light' if mode=='candidate' else 'muted',1.1,True,True)
    for u,v in RES:edge(u,v,'navy' if mode!='raw' else 'muted',2 if mode!='raw' else 1.3,mode!='raw',mode=='final' and u==v)
    for u,v in TREE:edge(u,v,'ink',1.9 if mode!='raw' else 1.6,False,mode!='raw')
    ax,ay=pts['account']
    cc.circle(ax-.055*scale,ay-.69*scale,.11*scale,'paper','muted',.8)
    cc.line(ax,ay-.58*scale,ax,ay-rad,'muted',1,True,True)
    cc.text('entry',ax-.28*scale,ay-.94*scale,.56*scale,.20*scale,10,align=PP_ALIGN.CENTER)
    for k,(xx,yy) in pts.items():
        cc.circle(xx-rad,yy-rad,rad*2,'cream' if mode=='final' else ('purple' if k=='update' else 'pale_blue'),'ink',1.2)
        ss=.34*scale;cc.icon(k,xx-.46*ss,yy-.41*ss,ss,'cream' if k=='reason' else 'pale_blue')
        cc.text(LABEL[k],xx-.47*scale,yy+rad+.04*scale,.94*scale,.26*scale,12.8,True,PP_ALIGN.CENTER)
    return pts


raw=graph(c,.65,2.75,'raw',1.0)
# Two small evidence labels; these are symbols, not made-up experimental scores.
c.rect(1.44,3.46,.53,.24,'white',None)
c.text('nₑ',1.44,3.45,.53,.26,12,True,PP_ALIGN.CENTER)
c.rect(2.54,3.10,.48,.25,'white',None)
c.text('wₑ',2.54,3.10,.48,.25,12,True,PP_ALIGN.CENTER)
note(c,'context stays attached',.77,5.49,3.08,.31,13,angle=-3)
c.icon('chat',4.00,5.52,.36,'pale_blue')
c.path([(3.91,5.61),(3.63,5.61),(3.43,5.27)],'muted',1.0,True,True)

# Main flow arrows have a small angle rather than the former boxed grid.
a=c.arrow(5.42,3.64,.56,.38);a.rotation=-4
a=c.arrow(11.79,3.64,.56,.38);a.rotation=4

candidate=graph(c,6.69,2.52,'candidate',1.0)
highlighter(c,8.00,1.51,.90)
note(c,'shared spine',9.27,1.65,2.12,.36,15,angle=-4)
c.path([(9.52,2.09),(9.15,2.19),(8.91,2.46)],'muted',1,True)
scissors(c,10.86,3.08,.82)
note(c,'trim / swap',10.66,5.56,1.56,.30,13,angle=3)
swap_icon(c,11.13,5.04,.72)
zoom(c,6.63,4.53,.70)
note(c,'check evidence',7.38,5.49,2.64,.28,13,angle=-3)

final=graph(c,12.91,2.55,'final',1.0)
note(c,'a graph, with branches & retries',12.95,1.60,4.81,.40,15,angle=-2)
c.check(17.64,1.55,.44)
note(c,'alternative',16.16,3.52,1.53,.28,12,angle=3)
note(c,'retry',14.98,5.41,1.15,.30,12,angle=-3)
c.math(['G = B ∪ E',('R',-24000)],16.39,5.36,1.71,.41,21)

# The objective is an annotated strip, not a large formula box.
marker(c,[(.55,6.15),(18.03,6.12)],'neon',25,27000)
c.math(['min',('G',-24000),'  J(G) = C(G) − βW',('tree',-24000),'(G)'],.76,5.98,7.31,.47,24)
c.math(['s.t.  F',('ρ',-24000),'(G) ≥ η  and observed-entry reachability'],8.31,6.02,9.47,.38,18)

# Three visual explanations below the highlighted objective.
# Subtle paper backgrounds have no surrounding stroke.
c.rect(.42,6.79,5.60,2.16,'cream',None)
c.rect(6.30,6.79,5.83,2.16,'paper',None)
c.rect(12.41,6.79,5.75,2.16,'pale_teal',None)

c.icon('chat',.66,7.01,.43,'pale_blue')
c.text('Cover diverse traces',1.23,6.95,4.36,.35,19,True)
c.text('ρ = 0.75',4.46,7.41,1.20,.23,12,align=PP_ALIGN.RIGHT)
c.text('retained → capped score',2.64,7.41,1.83,.23,10.7,align=PP_ALIGN.RIGHT)
for i,k in enumerate([4,3,2]):
    yy=7.81+i*.31
    c.text(f'τ{i+1}',.73,yy-.04,.44,.23,13,True)
    # Four edge occurrences, with dots marking their boundaries.
    for j in range(4):
        xx=1.35+j*.45
        marker(c,[(xx,yy+.07),(xx+.40,yy+.07)],'mint' if j<k else 'gray',9,90000)
        c.line(xx,yy+.07,xx+.40,yy+.07,'ink' if j<k else 'light',1.0)
        c.circle(xx-.025,yy+.045,.05,'ink' if j<k else 'muted',None)
    c.circle(3.10,yy+.045,.05,'ink',None)
    c.text(f'{k/4:.2f}  →  {min((k/4)/.75,1):.2f}',3.42,yy-.04,1.53,.23,13.5,True,PP_ALIGN.CENTER)
    if k>=3:c.check(5.16,yy-.07,.26)
    else:
        c.circle(5.16,yy-.07,.26,'white','ink',.7)
        c.text('⅔',5.16,yy-.06,.26,.23,10,align=PP_ALIGN.CENTER)
c.text('Cap each trace at 1, then average.',.76,8.74,4.91,.20,11.5,align=PP_ALIGN.CENTER)

c.icon('graph',6.55,7.01,.44)
c.text('Compact, supported structure',7.13,6.95,4.73,.35,18.5,True)
# Cost picture: five competing successors, followed by pruning to three.
fanx,fany=6.94,7.98
for i in range(5):
    yy=7.57+i*.23
    c.line(fanx,fany,7.77,yy,'light' if i in [0,4] else 'ink',1.1)
    c.circle(7.69,yy-.075,.15,'paper' if i in [0,4] else 'pale_blue','ink',.65)
c.circle(fanx-.11,fany-.11,.22,'yellow','ink',.9)
scissors(c,7.97,7.67,.48)
c.text('C(G) ↓',6.57,8.66,2.26,.23,15,True,PP_ALIGN.CENTER)
# Strong backbone picture: a small highlighted rooted tree.
for a,b in [((9.32,7.94),(9.98,7.61)),((9.32,7.94),(9.98,8.29)),((9.98,7.61),(10.71,7.61))]:
    marker(c,[a,b],'neon',13,65000);c.line(*a,*b,'ink',1.2,True)
for xx,yy in [(9.32,7.94),(9.98,7.61),(9.98,8.29),(10.71,7.61)]:
    c.circle(xx-.10,yy-.10,.20,'pale_blue','ink',.9)
zoom(c,11.12,7.91,.41)
c.text('Wtree(G) ↑',9.14,8.66,2.65,.23,15,True,PP_ALIGN.CENTER)

c.check(12.68,7.03,.34)
c.text('Accept feasible improvements',13.17,6.95,4.68,.35,18.5,True)
for i,text in enumerate(['Coverage ≥ η','All actions remain entry-reachable','Joint objective strictly decreases']):
    yy=7.49+i*.34
    c.check(12.75,yy+.02,.25)
    c.text(text,13.17,yy,4.48,.28,13.4)
swap_icon(c,12.75,8.60,.47)
c.text('Bounds → shortlist → recompute B',13.39,8.63,4.33,.26,12.4)

# A concise implementation cue connects the formal objective to the search.
c.text('Approximation:',.51,9.15,1.68,.23,12,True)
c.text('tree seed → coverage additions → delete / one-edge exchange → bound screening + backbone recomputation',2.15,9.15,14.25,.23,12)
c.text('Schematic',16.59,9.15,1.44,.23,11.5,align=PP_ALIGN.RIGHT)

slide.notes_slide.notes_text_frame.text=(
    'TraceCover detail, marker annotation style v2. All marks, pictograms, graph nodes and '
    'formula text are editable native PPT artwork. Yellow highlighter denotes the selected '
    'backbone; blue dashed arrows retain alternative transitions and retries. Scissors and '
    'swap pictograms denote deletion and replacement proposals, not semantic decisions or '
    'LLM interventions. Magnifiers denote support/specificity weighting and candidate review. '
    'The weighted graph and its retained subgraph are schematic, not an exported ABCD SOP. '
    'The virtual entry is non-executable and attaches to an observed starting action. '
    'The coverage example is independent of the graph: 4/4, 3/4, 2/4 occurrence coverage, '
    'rho=.75, capped scores 1,1,2/3, averaged over dialogues. '
    'All nodes remain; the selected graph contains a spanning backbone and non-tree edges. '
    'The maximum arborescence is recomputed inside each shortlisted candidate, so the seed '
    'backbone is not immutable. The final graph passes to evidence induction and text '
    'compilation, not directly to the agent. Search is budgeted, without an optimality guarantee.')

prs.save(PPTX)
check=Presentation(PPTX)
assert len(check.slides)==1
def audit(shapes):
    n=0
    for sh in shapes:
        assert sh.shape_type!=13
        if sh.shape_type==6:n+=audit(sh.shapes)
        else:n+=1
    return n
n=audit(check.slides[0].shapes)
for sh in check.slides[0].shapes:
    assert sh.left>=0 and sh.top>=0,sh.name
    assert sh.left+sh.width<=check.slide_width+100,(sh.name,'right')
    assert sh.top+sh.height<=check.slide_height+100,(sh.name,'bottom')

manifest={'version':2,'style':'marker annotations, editable diagrams, no rigid framing',
          'native_objects':n,'raster_objects':0,'backbone_edges':TREE,'residual_edges':RES,
          'illustrative':True,'coverage':{'rho':.75,'retained_occurrences':[4,3,2],'total_occurrences':4}}
(OUT/'figure_manifest_v2.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(f'Created {PPTX.name}; {n} editable objects; one slide.')


def export():
    import fitz
    import win32com.client
    app=win32com.client.DispatchEx('PowerPoint.Application');pres=None
    try:
        pres=app.Presentations.Open(str(PPTX),ReadOnly=True,Untitled=False,WithWindow=False)
        pres.SaveAs(str(OUT/'TraceCover_detail_v2.pdf'),32)
        pres.Slides(1).Export(str(OUT/'TraceCover_detail_v2.png'),'PNG',3348,1710)
        with fitz.open(OUT/'TraceCover_detail_v2.pdf') as doc:
            assert len(doc)==1
            assert 'TraceCover' in doc[0].get_text()
            print('PDF:',len(doc[0].get_drawings()),'vector paths;',len(doc[0].get_images()),'embedded images.')
        print('Exported PDF and PNG from actual PowerPoint rendering.')
    finally:
        if pres is not None:pres.Close()
        app.Quit()


if __name__=='__main__':export()
