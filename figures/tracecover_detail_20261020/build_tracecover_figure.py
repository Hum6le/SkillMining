"""TraceCover detail figure, built from editable native PowerPoint shapes.

Run from any directory. Microsoft PowerPoint renders the PPTX to vector PDF
and PNG. The graphs and coverage bars are explicitly illustrative, not results.
"""
from __future__ import annotations

import copy
import json
import math
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '.pptx_deps'))
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent
PPTX = OUT / 'TraceCover_detail_editable_v1.pptx'
W, H = 18.6, 10.0
C = dict(ink='182235', black='111111', blue='42B9ED', navy='4678C4',
         pale_blue='DDEDF9', cream='FFF3CB', yellow='FFD75F', coral='F36F63',
         teal='59CCAF', pale_teal='DEF4EA', purple='E5DAF3', gray='E3E5E8',
         paper='F7F8FA', muted='667385', green='339866', red='D65555',
         light='C9CFD7', white='FFFFFF')


def rgb(value):
    return RGBColor.from_string(C.get(value, value))


class Canvas:
    def __init__(self, parent, ox=0., oy=0.):
        self.parent, self.ox, self.oy = parent, ox, oy

    def shape(self, kind, x, y, w, h, fill='white', line='ink', lw=1.3):
        s = self.parent.add_shape(kind, Inches(self.ox+x), Inches(self.oy+y), Inches(w), Inches(h))
        if fill is None:
            s.fill.background()
        else:
            s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = .09
        s._element.spPr.append(OxmlElement('a:effectLst'))
        for ref in s._element.xpath('./p:style/a:effectRef'):
            ref.set('idx', '0')
        return s

    def rect(self, x, y, w, h, fill='white', line='ink', lw=1.1, rounded=True):
        return self.shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                          x,y,w,h,fill,line,lw)

    def circle(self, x,y,d,fill='white',line='ink',lw=1.2):
        return self.shape(MSO_SHAPE.OVAL,x,y,d,d,fill,line,lw)

    def text(self, value,x,y,w,h,size=17,bold=False,align=PP_ALIGN.LEFT,
             color='black',font='Arial'):
        s=self.parent.add_textbox(Inches(self.ox+x),Inches(self.oy+y),Inches(w),Inches(h))
        tf=s.text_frame;tf.clear();tf.word_wrap=True
        tf.margin_left=tf.margin_right=Inches(.025)
        tf.margin_top=tf.margin_bottom=Inches(.01)
        tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        for i,line in enumerate(str(value).split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.alignment=align;p.space_before=p.space_after=Pt(0);p.line_spacing=1.03
            r=p.add_run();r.text=line;r.font.name=font;r.font.size=Pt(size)
            r.font.bold=bold;r.font.color.rgb=rgb(color)
        return s

    def math(self, parts,x,y,w,h,size=19,align=PP_ALIGN.CENTER):
        s=self.text('',x,y,w,h,size,align=align,font='Cambria Math')
        p=s.text_frame.paragraphs[0]
        for part in parts:
            value,shift = part if isinstance(part,tuple) else (part,0)
            r=p.add_run();r.text=value;r.font.name='Cambria Math'
            r.font.size=Pt(size if not shift else size*.72)
            r.font.color.rgb=rgb('black')
            if shift:r._r.get_or_add_rPr().set('baseline',str(shift))
        return s

    def line(self,x1,y1,x2,y2,color='ink',width=1.5,arrow=False,dash=False):
        s=self.parent.add_connector(MSO_CONNECTOR.STRAIGHT,
            Inches(self.ox+x1),Inches(self.oy+y1),Inches(self.ox+x2),Inches(self.oy+y2))
        s.line.color.rgb=rgb(color);s.line.width=Pt(width)
        if dash:s.line.dash_style=MSO_LINE_DASH_STYLE.DASH
        if arrow:
            tail=OxmlElement('a:tailEnd');tail.set('type','triangle');tail.set('w','med');tail.set('len','med')
            s.line._get_or_add_ln().append(tail)
        s._element.spPr.append(OxmlElement('a:effectLst'))
        for ref in s._element.xpath('./p:style/a:effectRef'):
            ref.set('idx','0')
        return s

    def path(self,points,color='ink',width=1.6,arrow=True,dash=False):
        # A single editable freeform retains an arrowhead only at its endpoint.
        pts=[(int(Inches(self.ox+x)),int(Inches(self.oy+y))) for x,y in points]
        fb=self.parent.build_freeform(*pts[0]);fb.add_line_segments(pts[1:],close=False)
        s=fb.convert_to_shape();s.fill.background()
        s.line.color.rgb=rgb(color);s.line.width=Pt(width)
        if dash:s.line.dash_style=MSO_LINE_DASH_STYLE.DASH
        if arrow:
            tail=OxmlElement('a:tailEnd');tail.set('type','triangle');tail.set('w','med');tail.set('len','med')
            s.line._get_or_add_ln().append(tail)
        s._element.spPr.append(OxmlElement('a:effectLst'))
        for ref in s._element.xpath('./p:style/a:effectRef'):
            ref.set('idx','0')
        return s

    def arrow(self,x,y,w=.32,h=.27,direction='right',fill='coral'):
        return self.shape({'right':MSO_SHAPE.RIGHT_ARROW,'down':MSO_SHAPE.DOWN_ARROW,
                           'up':MSO_SHAPE.UP_ARROW}[direction],x,y,w,h,fill,'ink',.9)

    def check(self,x,y,s=.27):
        self.circle(x,y,s,'pale_teal','green',.8)
        self.line(x+.22*s,y+.51*s,x+.43*s,y+.72*s,'ink',1.5)
        self.line(x+.43*s,y+.72*s,x+.79*s,y+.27*s,'ink',1.5)

    def icon(self,kind,x,y,s=.5,fill='pale_blue'):
        if kind=='account':
            self.rect(x,y,s,.75*s,fill,'ink',1)
            self.circle(x+.11*s,y+.10*s,.24*s,'yellow','ink',.8)
            self.shape(MSO_SHAPE.ARC,x+.08*s,y+.35*s,.31*s,.22*s,'navy','ink',.7)
            for j in range(3):self.line(x+.50*s,y+(.18+.15*j)*s,x+.86*s,y+(.18+.15*j)*s,width=.9)
        elif kind=='validate':
            self.circle(x,y,s,'teal','ink',1.1)
            self.line(x+.22*s,y+.53*s,x+.44*s,y+.74*s,width=1.8)
            self.line(x+.44*s,y+.74*s,x+.80*s,y+.26*s,width=1.8)
        elif kind=='status':
            self.rect(x,y+.16*s,.61*s,.44*s,'yellow','ink',.8,False)
            self.rect(x+.61*s,y+.29*s,.29*s,.31*s,'coral','ink',.8,False)
            self.circle(x+.11*s,y+.52*s,.20*s,'ink','ink',.5)
            self.circle(x+.65*s,y+.52*s,.20*s,'ink','ink',.5)
        elif kind=='refund':
            self.circle(x,y,s,'yellow','ink',1)
            self.circle(x+.12*s,y+.12*s,.76*s,'cream','ink',.6)
            self.text('$',x+.1*s,y+.04*s,.8*s,.9*s,max(10,s*39),True,PP_ALIGN.CENTER)
        elif kind in {'reason','update','doc'}:
            self.shape(MSO_SHAPE.FOLDED_CORNER,x,y,.76*s,s,fill,'ink',.9)
            for j in range(3):self.line(x+.13*s,y+(.30+.16*j)*s,x+.59*s,y+(.30+.16*j)*s,width=.8)
            if kind=='update':
                p=self.rect(x+.53*s,y+.30*s,.15*s,.67*s,'yellow','ink',.6,False);p.rotation=35
        elif kind=='chat':
            self.shape(MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,x,y,s,.7*s,fill,'ink',.9)
            for j in range(3):self.circle(x+(.18+.23*j)*s,y+.23*s,.08*s,'ink','ink',.4)
        elif kind=='search':
            self.circle(x,y,.65*s,None,'ink',1.5)
            self.line(x+.55*s,y+.54*s,x+.95*s,y+.94*s,width=2)
        elif kind=='graph':
            for a,b in [((.16,.17),(.77,.20)),((.16,.17),(.47,.80)),((.77,.20),(.47,.80))]:
                self.line(x+a[0]*s,y+a[1]*s,x+b[0]*s,y+b[1]*s,width=1)
            for a,b in [(.04,.04),(.65,.07),(.35,.67)]:self.circle(x+a*s,y+b*s,.25*s,'blue','ink',.8)


prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
slide=prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb('white')
c=Canvas(slide.shapes)
c.text('TraceCover',.30,.16,3.7,.49,30,True)
c.text('Joint selection of a compact procedural graph and its evidence-weighted backbone',3.34,.21,14.93,.37,22)
c.text('Illustrative graph and coverage scores · offline, training traces only',.33,.71,10.2,.24,13)

# Shared legend: line type is meaningful independently of color.
for xx,color,dash,label,ww in [(11.12,'coral',False,'Backbone B',2.08),
                              (13.63,'navy',True,'Retained residual',2.5),
                              (16.36,'light',True,'Not selected',1.96)]:
    c.line(xx,.79,xx+.28,.79,color,2.1,True,dash)
    c.text(label,xx+.35,.64,ww-.35,.29,12.5)

PANELS={}
def panel(key,x,w,title,fill):
    group=slide.shapes.add_group_shape();group.name=key
    cc=Canvas(group.shapes,x,1.12)
    cc.rect(0,0,w,6.04,'white','ink',1.2)
    cc.rect(.08,.08,w-.16,.53,fill,'ink',.7)
    cc.text(title,.20,.14,w-.40,.36,19.5,True)
    PANELS[key]=group
    return cc

left=panel('a_observed_graph',.26,4.72,'(a) Observed action transitions','cream')
mid=panel('b_joint_criterion',5.40,7.34,'(b) Three coupled selection requirements','cream')
right=panel('c_selected_graph',13.16,5.17,'(c) Selected procedural structure','pale_blue')
c.arrow(5.02,3.38,.29,.29)
c.arrow(12.80,3.38,.29,.29)

NODES={'account':(.48,1.19),'validate':(1.46,1.19),
       'status':(2.64,.42),'refund':(3.87,.42),
       'reason':(2.64,2.02),'update':(3.87,2.02)}
TREE=[('account','validate'),('validate','status'),('status','refund'),
      ('validate','reason'),('reason','update')]
RESIDUAL=[('status','update'),('reason','reason')]
OMITTED=[('refund','update'),('update','refund')]
LABEL={'account':'Account','validate':'Validate','status':'Status','refund':'Refund',
       'reason':'Reason','update':'Update'}


def action_graph(cc,ox,oy,scale=1.,selected=False):
    pts={k:(ox+x*scale,oy+y*scale) for k,(x,y) in NODES.items()}
    rad=.25*scale
    def ordinary(edge,color,width,dash=False):
        u,v=edge;x1,y1=pts[u];x2,y2=pts[v]
        if u==v:
            cc.path([(x1-.16*scale,y1+.20*scale),(x1-.38*scale,y1+.40*scale),
                     (x1-.38*scale,y1+.78*scale),(x1+.40*scale,y1+.78*scale),
                     (x1+.40*scale,y1+.40*scale),(x1+.16*scale,y1+.20*scale)],
                    color,width,True,dash)
        else:
            dx,dy=x2-x1,y2-y1;d=math.hypot(dx,dy)
            cc.line(x1+rad*dx/d,y1+rad*dy/d,x2-rad*dx/d,y2-rad*dy/d,color,width,True,dash)
    for edge in OMITTED:
        # Parallel observed directions are offset to keep their arrowheads legible.
        u,v=edge;x1,y1=pts[u];x2,y2=pts[v]
        offset=(.31 if u=='refund' else .52)*scale
        xx=x1+offset
        cc.path([(x1+rad*.8,y1+(rad*.45 if u=='refund' else -rad*.45)),
                 (xx,y1+(rad*.45 if u=='refund' else -rad*.45)),
                 (xx,y2+(-rad*.45 if u=='refund' else rad*.45)),
                 (x2+rad*.8,y2+(-rad*.45 if u=='refund' else rad*.45))],
                 'light' if selected else 'muted',1.0 if selected else 1.2,True,True)
    for edge in RESIDUAL:ordinary(edge,'navy' if selected else 'muted',2 if selected else 1.4,selected)
    for edge in TREE:ordinary(edge,'coral' if selected else 'ink',2.3 if selected else 1.7)
    # Virtual root is organizational, not an executable service action.
    ax,ay=pts['account']
    cc.line(ax,ay-.55*scale,ax,ay-rad,'muted',1.0,True,True)
    cc.circle(ax-.06*scale,ay-.68*scale,.12*scale,'paper','muted',.7)
    cc.text('root',ax-.26*scale,ay-.99*scale,.53*scale,.25*scale,10.5,align=PP_ALIGN.CENTER)
    for name,(xx,yy) in pts.items():
        cc.circle(xx-rad,yy-rad,rad*2,'pale_blue' if name!='update' else 'purple','ink',1.1)
        s=.32*scale
        cc.icon(name,xx-s*.46,yy-s*.40,s,'cream' if name=='reason' else 'pale_blue')
        cc.text(LABEL[name],xx-.45*scale,yy+rad+.045*scale,.90*scale,.27*scale,12.5,True,PP_ALIGN.CENTER)
    return pts

# (a) Stacked trace cards preserve the visual vocabulary of the overview figure.
for j in [2,1,0]:
    left.rect(.32+j*.08,.91-j*.07,3.95,.75,'gray' if j else 'cream','ink',.8)
left.icon('chat',.47,1.07,.44,'pale_blue')
left.text('Training traces  D',1.02,.95,2.94,.29,17,True)
left.text('Account → Validate → Status → …',1.02,1.26,3.10,.22,12)
left.arrow(2.12,1.73,.25,.27,'down')
action_graph(left,.01,2.05,1.0,False)
left.math(['w',('e',-24000),' = log(1+n',('e',-24000),') + ½ log lift',('e',-24000)],.23,5.08,4.25,.31,17)
left.text('nₑ = distinct supporting dialogues',.30,5.43,4.12,.21,11.5,align=PP_ALIGN.CENTER)
left.rect(.25,5.71,4.22,.25,'pale_blue','ink',.7)
left.icon('chat',.36,5.73,.23)
left.text('Context · arguments · trace IDs stay linked',.75,5.72,3.57,.21,11.5)

# (b) Joint optimization objective, with the terms expanded immediately below.
mid.rect(.20,.78,6.94,.87,'cream','ink',1)
mid.math(['min',('G',-23000),'  J(G) = C(G) − β W',('tree',-24000),'(G)'],.32,.87,6.70,.39,25)
mid.math(['s.t.  F',('ρ',-23000),'(G) ≥ η;  observed edges and root reachability'],.34,1.29,6.66,.23,15)

mid.rect(.20,1.84,6.94,2.20,'pale_blue','ink',.85)
mid.text('TRACE COVERAGE',.37,1.99,4.70,.27,15,True)
mid.text('Illustration: ρ = 0.75',4.83,1.99,2.10,.27,12.5,align=PP_ALIGN.RIGHT)
mid.math(['r',('i',-24000),'(G) = retained transition occurrences / all occurrences'],.39,2.37,6.50,.28,16)
barx,bary,barw=1.03,2.91,2.75
mid.text('trace',.37,2.67,.58,.20,10.5)
mid.text('retained fraction rᵢ',barx,2.67,barw,.20,10.5,align=PP_ALIGN.CENTER)
mid.text('capped score',4.00,2.67,1.42,.20,10.5,align=PP_ALIGN.CENTER)
for i,r in enumerate([1.,.75,.5]):
    yy=bary+i*.29
    mid.text(f'τ{i+1}',.40,yy,.52,.23,13,True)
    mid.rect(barx,yy+.025,barw,.17,'white','ink',.5,False)
    mid.rect(barx,yy+.025,barw*r,.17,'teal','ink',.5,False)
    mid.text(f'{r:.2f}',3.91,yy,.55,.23,11.5,align=PP_ALIGN.RIGHT)
    mid.text(f'{min(r/.75,1):.2f}',4.79,yy,.58,.23,12.5,True,PP_ALIGN.CENTER)
mid.line(barx+barw*.75,2.85,barx+barw*.75,3.74,'ink',1.1,False,True)
mid.text('At most 1\nper trace',5.54,2.96,1.38,.48,13,True,align=PP_ALIGN.CENTER)
mid.math(['F',('ρ',-24000),'(G) = mean',('i',-24000),' min{1, r',('i',-24000),'(G)/ρ}'],.38,3.72,6.59,.23,15.5)

mid.rect(.20,4.13,3.36,1.70,'paper','ink',.85)
mid.text('STRUCTURAL COST',.35,4.28,3.05,.25,14.5,True)
mid.math(['C(G) = |E',('G',-24000),'| + γ ∑',('u',-24000),' p(u)'],.29,4.80,2.35,.31,15.5)
mid.text('(',2.63,4.62,.20,.60,36,font='Cambria Math',align=PP_ALIGN.CENTER)
numerator=mid.math(['d',('+',26000),('G',-24000),'(u)'],2.78,4.68,.58,.24,12)
numerator.text_frame.word_wrap=False
mid.math(['2'],2.78,4.98,.58,.22,12)
mid.text(')',3.36,4.62,.18,.60,36,font='Cambria Math',align=PP_ALIGN.CENTER)
mid.text('Edges + weighted successor pairs',.33,5.46,3.06,.23,12,align=PP_ALIGN.CENTER)

mid.rect(3.76,4.13,3.38,1.70,'pale_teal','ink',.85)
mid.text('BACKBONE EVIDENCE',3.91,4.28,3.04,.25,14.5,True)
mid.math(['W',('tree',-24000),'(G) = max',('B ∈ A(G)',-24000),' ∑',('e ∈ B',-24000),' w',('e',-24000)],3.88,4.70,3.14,.42,17)
mid.text('Best rooted spanning arborescence',3.90,5.18,3.06,.23,12,align=PP_ALIGN.CENTER)
mid.text('wₑ: support + transition specificity',3.90,5.48,3.06,.23,11.5,align=PP_ALIGN.CENTER)

# (c) Retained graph is a backbone plus residual transitions, including retries.
action_graph(right,.15,1.01,1.04,True)
right.text('alternative',3.38,2.07,1.14,.24,10.5,align=PP_ALIGN.CENTER)
right.text('retry',2.18,3.96,.73,.22,10.5,align=PP_ALIGN.CENTER)
right.math(['G = B ∪ E',('R',-24000)],.43,4.11,4.30,.43,24)
right.rect(.30,4.78,4.55,.83,'pale_blue','ink',.8)
right.icon('doc',.47,4.94,.46,'cream')
right.text('Compilation order + indexed evidence',1.06,4.91,3.59,.28,14,True)
right.text('Retained alternatives and retry links',1.06,5.26,3.59,.23,12)
right.text('Graph output → evidence induction → text skill',.33,5.72,4.48,.21,11.5,align=PP_ALIGN.CENTER)

# Practical approximation is visually separate from the formal criterion.
solver=slide.shapes.add_group_shape();solver.name='d_practical_approximation'
s=Canvas(solver.shapes,.26,7.43)
s.rect(0,0,18.07,2.32,'pale_blue','ink',1.15)
s.text('PRACTICAL APPROXIMATION',.20,.13,6.20,.29,17,True)
s.text('Training statistics fixed before search',11.95,.14,5.83,.25,12.5,align=PP_ALIGN.RIGHT)
sx=[.22,4.72,9.22,13.72]
titles=['1  Maximum-evidence tree','2  Coverage augmentation',
        '3  Local graph proposals','4  Bound screening + MaxTree']
subtitles=['Compute the initial backbone B₀',
           'Add residuals by gain ΔFρ / cost ΔC',
           'Delete e; or replace e with unused e′',
           'Recompute B inside each shortlisted graph']
for i,xx in enumerate(sx):
    s.rect(xx,.61,4.09,.94,'white','ink',.9)
    s.icon('graph' if i!=3 else 'search',xx+.14,.83,.40)
    s.text(titles[i],xx+.65,.72,3.32,.30,14.5,True)
    s.text(subtitles[i],xx+.65,1.12,3.32,.25,11.5)
    if i<3:s.arrow(xx+4.17,.95,.25,.25)
s.rect(.22,1.76,13.48,.37,'pale_teal','ink',.7)
s.check(.36,1.82,.24)
s.math(['Accept only if  F',('ρ',-24000),'(G′) ≥ η,  all actions remain root-reachable,  and  J(G′) < J(G)'],.78,1.79,12.67,.28,15)
s.text('Repeat within a bounded move budget',13.99,1.76,3.85,.38,12,align=PP_ALIGN.CENTER)
# Arrow from the solver to the selected graph; no optimality guarantee implied.
c.arrow(17.65,7.20,.27,.23,'up',fill='navy')

slide.notes_slide.notes_text_frame.text=(
    'TraceCover detail illustration. All objects are native PowerPoint shapes and editable. '
    'The left and right graphs are schematic, not a recovered ABCD SOP or a numerical mining result. '
    'A virtual root connects only observed dialogue entries; the illustrated root has one entry. '
    'Only service actions become graph nodes; raw context, arguments and trace IDs remain linked. '
    'The coverage bars form an independent three-trace numerical example with rho=0.75; they are not '
    'calculated from the schematic six-node graph. r_i is the fraction of transition OCCURRENCES retained; '
    'repeated transitions count multiplicity. Scores cap at one before averaging over dialogues. '
    'Graph cost is edge count plus source-frequency-weighted pairs of outgoing edges. '
    'The maximum rooted spanning arborescence is recomputed within each shortlisted candidate graph; '
    'the initial backbone is not immutable. Candidate deletions and one-edge exchanges must preserve '
    'coverage and observed-entry reachability and strictly improve the joint objective. '
    'The proposed search is budgeted and has no global optimality guarantee. '
    'Blue dashed edges are retained non-tree alternatives and retries; gray dashed edges are observed '
    'but unselected. The selected graph, not only its tree, passes to evidence induction and compilation.')

# Build two enlarged editable component sheets without re-running the generator.
def clone_group(source,dest,left,top,width,height):
    el=copy.deepcopy(source._element)
    dest.shapes._spTree.insert_element_before(el,'p:extLst')
    g=dest.shapes[-1]
    g.left=Inches(left);g.top=Inches(top);g.width=Inches(width);g.height=Inches(height)
    return g

for key,title,w0,h0 in [('b_joint_criterion','TraceCover · Selection criterion',7.34,6.04),
                       ('c_selected_graph','TraceCover · Output graph',5.17,6.04)]:
    detail=prs.slides.add_slide(prs.slide_layouts[6])
    dc=Canvas(detail.shapes)
    dc.text(title,.45,.24,17.2,.52,25,True)
    factor=1.42
    clone_group(PANELS[key],detail,(W-w0*factor)/2,.96,w0*factor,h0*factor)
    detail.notes_slide.notes_text_frame.text=slide.notes_slide.notes_text_frame.text

prs.save(PPTX)

# Meaningful integrity checks: graph/tree structure, numerical example, and deck geometry.
import networkx as nx
tree=nx.DiGraph(TREE);assert nx.is_arborescence(tree)
assert set(tree.nodes)==set(NODES)
assert set(TREE).isdisjoint(RESIDUAL)
assert ('reason','reason') in RESIDUAL
assert [min(v/.75,1.) for v in [1.,.75,.5]]==[1.,1.,2/3]
check=Presentation(PPTX);assert len(check.slides)==3
def audit_shapes(shapes):
    n=0
    for sh in shapes:
        assert sh.shape_type!=13,'No raster artwork should be embedded'
        if sh.shape_type==6:n+=audit_shapes(sh.shapes)
        else:n+=1
    return n
count=audit_shapes(check.slides[0].shapes)
for i,sl in enumerate(check.slides):
    for sh in sl.shapes:
        assert sh.left>=0 and sh.top>=0,(i,sh.name)
        assert sh.left+sh.width<=check.slide_width+100,(i,sh.name,'right')
        assert sh.top+sh.height<=check.slide_height+100,(i,sh.name,'bottom')

manifest={'figure':'TraceCover detail v1','canvas_inches':[W,H],'slides':3,
          'native_editable_objects_on_main_slide':count,'raster_objects':0,
          'source':'paper/overleaf_sync/sections/3_methodology.tex; skill_mining/joint_trace_cover.py',
          'illustration':'Schematic action graph; independent three-trace coverage example',
          'backbone_edges':TREE,'retained_residual_edges':RESIDUAL,'unselected_edges':OMITTED,
          'coverage_example':{'rho':.75,'retained_fractions':[1.,.75,.5],
                              'saturated_scores':[1.,1.,2/3]}}
(OUT/'figure_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(f'Created {PPTX.name}; {count} editable objects; 3 slides; no raster objects.')


def export():
    import fitz
    import win32com.client
    app=win32com.client.DispatchEx('PowerPoint.Application');pres=None
    try:
        pres=app.Presentations.Open(str(PPTX),ReadOnly=True,Untitled=False,WithWindow=False)
        working=OUT/'TraceCover_detail_working_pages_v1.pdf'
        pres.SaveAs(str(working),32)
        pres.Slides(1).Export(str(OUT/'TraceCover_detail_v1.png'),'PNG',3348,1800)
        for n in range(2,4):
            pres.Slides(n).Export(str(OUT/f'TraceCover_detail_component_{n-1}.png'),'PNG',2232,1200)
        with fitz.open(working) as doc:
            first=fitz.open();first.insert_pdf(doc,from_page=0,to_page=0)
            first.save(OUT/'TraceCover_detail_v1.pdf',garbage=4,deflate=True);first.close()
        with fitz.open(OUT/'TraceCover_detail_v1.pdf') as doc:
            assert len(doc)==1
            t=doc[0].get_text()
            assert 'TraceCover' in t and 'COVERAGE' in t and 'MaxTree' in t
            assert not doc[0].get_images(),'Main PDF should remain native vector/text'
        print('Exported actual PowerPoint rendering: single-page vector PDF + PNG + component sheets.')
    finally:
        if pres is not None:pres.Close()
        app.Quit()


if __name__=='__main__':
    export()
