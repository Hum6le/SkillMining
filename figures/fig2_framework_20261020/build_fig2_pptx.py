"""Editable Fig. 2: native PowerPoint shapes, with a complete overview and step sheets.

Run from the project root. PowerPoint COM exports the actual deck to PDF/PNG.
No generated full-slide raster is embedded in the presentation.
"""
from __future__ import annotations

import copy
import json
import math
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '.pptx_deps'))
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent
STEM = 'SkillMiner_Fig2_editable_v1'
PPTX = OUT / (STEM + '.pptx')
W, H = 20.0, 10.65
PW, PH = 4.55, 3.98
XS = [.35, 5.27, 10.19, 15.11]
YTOP, YBOT = .98, 6.16

C = dict(ink='182235', black='111111', blue='42B9ED', navy='4678C4',
         pale_blue='DDEDF9', cream='FFF3CB', yellow='FFD75F', coral='F36F63',
         teal='59CCAF', pale_teal='DEF4EA', purple='E5DAF3', gray='E3E5E8',
         paper='F7F8FA', muted='667385', green='339866', red='D65555', white='FFFFFF')

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor.from_string(C['white'])


def rgb(value):
    return RGBColor.from_string(C.get(value, value))


class Canvas:
    def __init__(self, parent, ox=0., oy=0.):
        self.parent, self.ox, self.oy = parent, ox, oy

    def shape(self, kind, x, y, w, h, fill='white', line='ink', lw=1.35, radius=.12):
        s = self.parent.add_shape(kind, Inches(self.ox+x), Inches(self.oy+y), Inches(w), Inches(h))
        if fill is None:
            s.fill.background()
        else:
            s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = radius
        # Override the presentation theme's default shape shadow.
        s._element.spPr.append(OxmlElement('a:effectLst'))
        for ref in s._element.xpath('./p:style/a:effectRef'):
            ref.set('idx','0')
        return s

    def rect(self, x, y, w, h, fill='white', line='ink', lw=1.35, rounded=True):
        return self.shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                          x,y,w,h,fill,line,lw)

    def circle(self, x, y, d, fill='white', line='ink', lw=1.35):
        return self.shape(MSO_SHAPE.OVAL,x,y,d,d,fill,line,lw)

    def text(self, value, x, y, w, h, size=18, bold=False, align=PP_ALIGN.LEFT,
             color='black', font='Arial', valign=MSO_ANCHOR.MIDDLE):
        s = self.parent.add_textbox(Inches(self.ox+x), Inches(self.oy+y), Inches(w), Inches(h))
        tf=s.text_frame; tf.clear(); tf.word_wrap=True
        tf.margin_left=tf.margin_right=Inches(.025)
        tf.margin_top=tf.margin_bottom=Inches(.015)
        tf.vertical_anchor=valign
        for i, line in enumerate(str(value).split('\n')):
            p=tf.paragraphs[0] if not i else tf.add_paragraph()
            p.alignment=align; p.space_before=Pt(0); p.space_after=Pt(0)
            p.line_spacing=1.02
            r=p.add_run();r.text=line;r.font.name=font;r.font.size=Pt(size)
            r.font.bold=bold;r.font.color.rgb=rgb(color)
        return s

    def line(self, x1,y1,x2,y2,color='ink',width=1.5,arrow=False,dash=False):
        s=self.parent.add_connector(MSO_CONNECTOR.STRAIGHT,
            Inches(self.ox+x1), Inches(self.oy+y1), Inches(self.ox+x2), Inches(self.oy+y2))
        s.line.color.rgb=rgb(color);s.line.width=Pt(width)
        if dash:s.line.dash_style=MSO_LINE_DASH_STYLE.DASH
        if arrow:
            tail=OxmlElement('a:tailEnd');tail.set('type','triangle');tail.set('w','med');tail.set('len','med')
            s.line._get_or_add_ln().append(tail)
        return s

    def arrow(self,x,y,w=.35,h=.26,direction='right'):
        kind={'right':MSO_SHAPE.RIGHT_ARROW,'left':MSO_SHAPE.LEFT_ARROW,
              'down':MSO_SHAPE.DOWN_ARROW}[direction]
        return self.shape(kind,x,y,w,h,'coral','ink',1.0)

    def picture(self,name,x,y,w,h=None):
        p=ROOT/'icons'/name
        return self.parent.add_picture(str(p),Inches(self.ox+x),Inches(self.oy+y),
            width=Inches(w),height=Inches(h) if h else None)

    def chip(self,label,x,y,w,h=.34,fill='pale_blue',size=14):
        self.rect(x,y,w,h,fill,'ink',.8)
        self.text(label,x+.03,y+.015,w-.06,h-.03,size,False,PP_ALIGN.CENTER)

    def icon(self,kind,x,y,s=.5,fill='pale_blue'):
        # All of these pictograms consist of native, editable PowerPoint shapes.
        if kind=='account':
            self.rect(x,y,s,s*.76,fill,'ink',1.2)
            self.circle(x+.10*s,y+.12*s,.23*s,'yellow','ink',.9)
            self.shape(MSO_SHAPE.ARC,x+.08*s,y+.35*s,.30*s,.23*s,'navy','ink',.7)
            for j in range(3):self.line(x+.51*s,y+(.20+.15*j)*s,x+.85*s,y+(.20+.15*j)*s,width=1)
        elif kind=='check':
            self.circle(x,y,s,'teal','ink',1.2)
            self.line(x+.22*s,y+.53*s,x+.43*s,y+.73*s,width=2)
            self.line(x+.43*s,y+.73*s,x+.80*s,y+.28*s,width=2)
        elif kind=='truck':
            self.rect(x,y+.16*s,s*.61,s*.44,'yellow','ink',1.0,False)
            self.rect(x+.61*s,y+.29*s,s*.29,s*.31,'coral','ink',1.0,False)
            self.circle(x+.11*s,y+.52*s,s*.20,'ink','ink',.6)
            self.circle(x+.65*s,y+.52*s,s*.20,'ink','ink',.6)
            self.line(x-.10*s,y+.30*s,x+.11*s,y+.30*s,width=1)
        elif kind=='coin':
            self.circle(x,y,s,'yellow','ink',1.2)
            self.circle(x+.12*s,y+.12*s,.76*s,'cream','ink',.65)
            self.text('$',x+.1*s,y+.06*s,.8*s,.88*s,max(10,s*40),True,PP_ALIGN.CENTER)
        elif kind in {'doc','edit','form','report'}:
            self.shape(MSO_SHAPE.FOLDED_CORNER,x,y,s*.76,s,fill,'ink',1.0)
            for j in range(3):self.line(x+.13*s,y+(.30+.16*j)*s,x+.59*s,y+(.30+.16*j)*s,width=.9)
            if kind=='edit':
                p=self.rect(x+.53*s,y+.33*s,.15*s,.67*s,'yellow','ink',.7,False)
                p.rotation=35
            if kind=='report':self.circle(x+.53*s,y+.68*s,.36*s,'cream','ink',.8)
        elif kind=='chat':
            self.shape(MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,x,y,s,s*.70,fill,'ink',1.0)
            for j in range(3):self.circle(x+(.19+.23*j)*s,y+.23*s,.08*s,'ink','ink',.5)
        elif kind=='search':
            self.circle(x,y,.66*s,None,'ink',1.6)
            self.line(x+.55*s,y+.54*s,x+.95*s,y+.94*s,width=2.2)
        elif kind=='graph':
            for p,q in [((.15,.22),(.72,.25)),((.15,.22),(.43,.78)),((.72,.25),(.43,.78))]:
                self.line(x+p[0]*s,y+p[1]*s,x+q[0]*s,y+q[1]*s,width=1)
            for a,b in [(.05,.1),(.6,.12),(.31,.65)]:self.circle(x+a*s,y+b*s,.25*s,'blue','ink',.8)
        elif kind=='cross':
            self.circle(x,y,s,'cream','red',1)
            self.line(x+.25*s,y+.25*s,x+.75*s,y+.75*s,'red',2)
            self.line(x+.25*s,y+.75*s,x+.75*s,y+.25*s,'red',2)
        elif kind=='clipboard':
            self.rect(x,y+.1*s,.76*s,.9*s,'cream','ink',1)
            self.rect(x+.20*s,y,.36*s,.18*s,'yellow','ink',.8)
            for j in range(3):self.line(x+.14*s,y+(.37+.18*j)*s,x+.61*s,y+(.37+.18*j)*s,width=.9)


full=Canvas(slide.shapes)
full.rect(.17,.14,19.66,4.99,'cream','ink',1.25)
full.rect(.17,5.37,19.66,5.09,'pale_blue','ink',1.25)
full.text('OFFLINE  ·  Structural prior induction',.42,.30,12,.42,24,True)
full.text('ONLINE  ·  Execution-guided refinement',.42,5.52,12,.42,24,True)
full.text('TraceCover',16.43,.30,2.9,.42,23,True,PP_ALIGN.RIGHT)
full.text('GRAFT',13.75,5.52,1.52,.40,23,True,PP_ALIGN.RIGHT)

TITLES={1:'Trace-to-graph',2:'TraceCover selection',3:'Evidence induction',4:'Skill compilation',
        5:'Agent rollouts',6:'Neighborhood sampling',7:'Contrastive diagnosis',8:'Consolidate & revise'}
NOTES={
1:'Project each training dialogue to canonical action transitions. Retain intervening dialogue context, ordered arguments, and trace IDs as linked evidence. The top layer is the observed action graph; the lower layer is raw trace evidence, not yet induced skill rules. The small graph is illustrative and is not a complete ABCD SOP.',
2:'Select a retained graph and its maximum-weight rooted arborescence under saturated trace coverage. The diagram separates backbone edges from retained alternative transitions. The solver uses an initializer followed by bounded feasible deletion/exchange search, without a global-optimality claim. Wtree is the maximum rooted-tree evidence weight.',
3:'Induce continuation conditions by comparing sibling-edge contexts. Obtain shared action rules and argument policies from grounded trace evidence. These documents are distinct from the raw evidence layer in step 1.',
4:'Compile an external skill entry and retrievable resources. The agent consumes text and resources; the graph remains an intermediate structure for provenance and subsequent refinement. K0=(T0,R0). No model parameter updates or agent training rollouts are needed for offline induction.',
5:'Run the fixed skill-guided agent on training prefixes. Compare predictions with training feedback. These are training rollouts, not held-out test execution. The diagram does not imply free-running dialogue simulation when the implementation predicts individual decisions from their prefixes.',
6:'Organize completed rollout cases using trajectory context, semantic similarity, graph locality and outcomes into contrastive batches. This schematic uses graph-neighborhood sampling to describe related-evidence selection; it makes no claim that only a reduced fraction of training turns is rolled out.',
7:'Diagnose local procedural issues using related successful/failed cases, relevant graph neighborhoods and skill-resource excerpts. Produce evidence-grounded reports covering routing conditions, operation requirements and argument binding.',
8:'Consolidate overlapping reports, reconcile proposed revisions against successful and failed evidence, and update skill text/resources and graph annotations. Process report groups with current skill state. Freeze the final skill for evaluation. No arrow loops back to fresh rollouts, and no regression guarantee is implied.'}

MODULES={}


def module(n):
    col=n-1 if n<=4 else 8-n
    x,y=XS[col],YTOP if n<=4 else YBOT
    g=slide.shapes.add_group_shape();g.name=f'Step_{n}_{TITLES[n].replace(" ","_")}'
    c=Canvas(g.shapes,x,y)
    c.rect(0,0,PW,PH,'white','ink',1.15)
    c.rect(.08,.08,PW-.16,.61,'cream' if n<=4 else 'pale_blue','ink',.75)
    c.text(f'{n}. {TITLES[n]}',.18,.13,PW-.36,.49,20,True)
    MODULES[n]=g
    return c


def mini_graph(c,x,y,scale=1.,selected=False,labels=False):
    pts={'A':(0,.47),'B':(.69,.47),'S':(1.40,.13),'F':(2.07,0),
         'D':(1.40,1.05),'U':(2.07,1.05)}
    edges=[('A','B'),('B','S'),('S','F'),('B','D'),('S','U'),('D','U')]
    for a,b in edges:
        p,q=pts[a],pts[b];dx=q[0]-p[0];dy=q[1]-p[1];L=math.hypot(dx,dy)
        off=.21
        c.line(x+(p[0]+off*dx/L)*scale,y+(p[1]+off*dy/L)*scale,
               x+(q[0]-off*dx/L)*scale,y+(q[1]-off*dy/L)*scale,
               'coral' if not selected or (a,b) != ('S','U') else 'navy',
               1.6,True,selected and (a,b) == ('S','U'))
    for k,(a,b) in pts.items():
        fill='cream' if k=='F' else ('purple' if k=='U' else 'pale_blue')
        c.circle(x+(a-.19)*scale,y+(b-.19)*scale,.38*scale,fill,'ink',1.0)
        if labels:
            c.text(k,x+(a-.17)*scale,y+(b-.15)*scale,.34*scale,.30*scale,14*scale,True,PP_ALIGN.CENTER)
        else:
            icon={'A':'account','B':'check','S':'truck','F':'coin','D':'clipboard','U':'edit'}[k]
            c.icon(icon,x+(a-.11)*scale,y+(b-.12)*scale,.23*scale)
    return {k:(x+a*scale,y+b*scale) for k,(a,b) in pts.items()}


def trace_card(c,x,y,w,label,utterance,alt=False):
    c.rect(x+.07,y-.07,w,.82,'gray','ink',.6)
    c.rect(x,y,w,.82,'pale_blue' if alt else 'cream','ink',1.0)
    c.picture('customer-female-633328.png' if alt else 'customer-apricot-F2BF80.png',x+.06,y+.07,.26)
    c.text(utterance,x+.35,y+.06,w-.41,.22,12,True)
    c.picture('customer-service.png',x+.06,y+.42,.28)
    for j,kind in enumerate(['account','check','edit' if alt else 'coin']):
        c.icon(kind,x+.42+j*.30,y+.44,.22)
        if j<2:c.line(x+.64+j*.30,y+.54,x+.70+j*.30,y+.54,'coral',1.1,True)
    c.text(label,x+w-.30,y+.23,.24,.19,11,True,PP_ALIGN.RIGHT)


c=module(1)
trace_card(c,.15,1.03,1.38,'τ₁','Delivered')
trace_card(c,.15,2.04,1.38,'τ₂','Order received',True)
c.rect(.27,3.04,1.22,.43,'cream','ink',.6)
c.text('…  τₙ',.32,3.08,1.1,.29,19,True,PP_ALIGN.CENTER)
c.arrow(1.58,2.0,.28,.26)
c.rect(1.94,.91,2.43,1.94,'gray','ink',.8)
c.text('Action graph',2.03,.99,2.20,.23,15,True)
nodes=mini_graph(c,2.22,1.64,.86)
c.line(2.44,2.77,2.29,3.13,'ink',.9,dash=True)
c.line(3.35,2.74,3.17,3.13,'ink',.9,dash=True)
c.line(3.89,2.72,4.02,3.13,'ink',.9,dash=True)
c.rect(1.94,3.03,2.43,.64,'pale_blue','ink',.8)
for j,(name,kind) in enumerate([('Ctx','chat'),('Args','form'),('IDs','account')]):
    x=2.04+.78*j;c.icon(kind,x,3.15,.23)
    c.text(name,x+.25,3.15,.45,.24,11,False)
c.text('Trace evidence',2.08,3.72,2.15,.18,12,False,PP_ALIGN.CENTER)

c=module(2)
mini_graph(c,.80,1.16,1.26,True)
c.text('Backbone',.37,2.84,1.10,.23,14,True)
c.line(.23,2.96,.34,2.96,'coral',2.1)
c.text('Alternatives',2.23,2.84,1.69,.23,14,True)
c.line(2.01,2.96,2.20,2.96,'navy',1.7,dash=True)
c.rect(.22,3.17,4.11,.61,'cream','ink',.8)
mathbox=c.text('',.28,3.20,3.99,.30,20,False,PP_ALIGN.CENTER,font='Cambria Math')
p=mathbox.text_frame.paragraphs[0]
for value,size,baseline in [('min  C(G) − β W',20,None),('tree',12,'-25000'),('(G)',20,None)]:
    r=p.add_run();r.text=value;r.font.name='Cambria Math';r.font.size=Pt(size);r.font.color.rgb=rgb('black')
    if baseline:r._r.get_or_add_rPr().set('baseline',baseline)
mathbox=c.text('',.28,3.51,3.99,.23,17,False,PP_ALIGN.CENTER,font='Cambria Math')
p=mathbox.text_frame.paragraphs[0]
for value,size,baseline in [('s.t.  F',17,None),('ρ',11,'-25000'),('(G) ≥ η',17,None)]:
    r=p.add_run();r.text=value;r.font.name='Cambria Math';r.font.size=Pt(size);r.font.color.rgb=rgb('black')
    if baseline:r._r.get_or_add_rPr().set('baseline',baseline)

c=module(3)
for j in range(3):c.icon('doc',.36+j*.065,1.09+j*.06,.69,'cream')
c.icon('chat',.38,2.05,.66,'pale_blue')
c.text('Grounded\ntrace evidence',.16,2.88,1.25,.55,15,True,PP_ALIGN.CENTER)
c.arrow(1.23,1.97,.32,.29)
for j,(title,sub,kind,fill) in enumerate([
    ('Branch conditions','Context → continuation','chat','cream'),
    ('Action rules','Operation requirements','clipboard','pale_blue'),
    ('Argument policies','Value sources & binding','form','pale_teal')]):
    yy=.98+j*.86;c.rect(1.68,yy,2.65,.70,fill,'ink',.95)
    c.icon(kind,1.80,yy+.15,.36)
    c.text(title,2.25,yy+.05,1.98,.29,15,True)
    c.text(sub,2.25,yy+.33,1.98,.26,12)
c.text('Compare sibling transitions',1.65,3.63,2.69,.22,14,False,PP_ALIGN.CENTER)

c=module(4)
c.rect(.35,.99,3.85,2.21,'gray','ink',1.1)
c.rect(.62,1.20,3.31,.65,'pale_blue','ink',1.0)
c.icon('doc',.76,1.31,.35)
c.text('skill.md',1.27,1.25,2.39,.29,22,True)
c.text('Procedure + resource pointers',.57,1.91,3.47,.27,15,False,PP_ALIGN.CENTER)
for j,(t,fill) in enumerate([('Action\nrules','cream'),('Argument\npolicies','pale_teal'),('Transition\nevidence','pale_blue')]):
    xx=.59+j*1.16;c.rect(xx,2.35,1.04,.57,fill,'ink',.8)
    c.text(t,xx+.02,2.37,1.0,.49,13,False,PP_ALIGN.CENTER)
    c.line(2.25,2.17,xx+.52,2.33,'navy',.8,False,True)
c.text('Load detailed resources on demand',.34,3.32,3.87,.30,16,False,PP_ALIGN.CENTER)
c.text('K⁰ = (T⁰, R⁰)',.40,3.66,3.75,.25,21,True,PP_ALIGN.CENTER,font='Cambria Math')

c=module(5)
for j in range(2):c.rect(.28+j*.06,1.03-j*.06,1.25,.77,'cream','ink',.75)
c.text('Training\nprefixes',.37,1.12,1.03,.51,17,True,PP_ALIGN.CENTER)
c.icon('doc',3.34,1.0,.55,'pale_blue')
c.text('Skill K⁰',2.92,1.60,1.17,.23,16,True,PP_ALIGN.CENTER)
c.picture('customer-service.png',1.88,1.55,.77)
c.line(1.55,1.75,1.87,1.94,'coral',1.5,True)
c.line(3.24,1.78,2.66,1.94,'coral',1.5,True)
c.text('Fixed agent',1.32,2.35,1.90,.27,16,True,PP_ALIGN.CENTER)
c.line(2.27,2.66,2.27,2.92,'coral',1.5,True)
c.chip('Predictions',.44,3.02,1.40,.44,'pale_blue',17)
c.chip('Feedback',2.73,3.02,1.37,.44,'cream',17)
c.icon('check',1.95,3.07,.26);c.icon('cross',2.30,3.07,.26)
c.text('Training executions only',.40,3.62,3.75,.25,15,False,PP_ALIGN.CENTER)

c=module(6)
local=[(.54,1.4),(1.13,1.08),(1.48,1.78),(.65,2.23),(1.85,2.40)]
for a,b in [(0,1),(0,3),(1,2),(2,3),(2,4)]:
    p,q=local[a],local[b];c.line(*p,*q,'muted',1.0)
for j,(x,y) in enumerate(local):c.circle(x-.14,y-.14,.28,'yellow' if j==2 else 'pale_blue','ink',.8)
s=c.rect(.28,.87,1.48,1.66,None,'navy',1.0);s.line.dash_style=MSO_LINE_DASH_STYLE.DASH
c.text('Local graph',.22,2.65,1.62,.27,15,True,PP_ALIGN.CENTER)
c.arrow(1.89,1.68,.35,.27)
for j in range(3):
    xx=2.43+(j%2)*.96;yy=1.0+(j//2)*.83
    c.rect(xx,yy,.85,.67,'pale_teal' if j!=1 else 'cream','ink',.8)
    c.icon('cross' if j==1 else 'check',xx+.08,yy+.10,.23)
    for z in range(2):c.line(xx+.38,yy+.22+z*.16,xx+.72,yy+.22+z*.16,'muted',.7)
c.text('Rollout cases',2.24,2.65,2.06,.27,15,True,PP_ALIGN.CENTER)
for j in range(3):c.chip(f'B{j+1}',.76+j*1.07,3.12,.83,.44,'pale_blue',19)
c.text('Contrastive evidence batches',.36,3.64,3.84,.23,15,False,PP_ALIGN.CENTER)

c=module(7)
c.rect(.27,1.0,1.20,.93,'pale_teal','ink',.8)
c.icon('check',.40,1.13,.28);c.icon('cross',.87,1.13,.28)
c.text('Cases',.35,1.58,1.03,.24,16,True,PP_ALIGN.CENTER)
c.rect(2.92,1.0,1.32,.93,'cream','ink',.8)
c.icon('doc',3.07,1.14,.36);c.icon('graph',3.62,1.14,.38)
c.text('Resources',2.99,1.58,1.18,.24,13,True,PP_ALIGN.CENTER)
c.line(1.51,1.75,1.91,2.23,'coral',1.5,True)
c.line(3.00,1.75,2.50,2.23,'coral',1.5,True)
c.icon('search',1.94,2.20,.67)
c.line(2.74,2.53,3.02,2.53,'coral',1.5,True)
c.icon('report',3.21,2.12,.66,'pale_blue')
c.text('Report dⱼ',2.95,2.97,1.26,.27,17,True,PP_ALIGN.CENTER)
for j,label in enumerate(['Conditions','Operations','Binding']):
    c.chip(label,.25+1.38*j,3.48,1.28,.32,'pale_blue',12)

c=module(8)
for j in range(3):c.icon('report',.32+j*.09,1.02-j*.05,.64,'pale_blue')
c.text('Reports',.25,1.79,1.10,.27,17,True,PP_ALIGN.CENTER)
c.arrow(1.31,1.29,.34,.26)
c.rect(1.78,1.10,1.31,.71,'cream','ink',.95)
c.text('Reconcile',1.80,1.27,1.27,.30,16,True,PP_ALIGN.CENTER)
c.arrow(3.23,1.29,.31,.26)
c.icon('edit',3.65,1.06,.56,'pale_teal')
c.text('Updates',3.40,1.79,.99,.27,17,True,PP_ALIGN.CENTER)
c.line(3.88,2.11,3.88,2.35,'coral',1.5,True)
c.rect(.39,2.48,3.78,.75,'gray','ink',1.0)
c.icon('graph',.56,2.62,.44);c.icon('doc',1.17,2.60,.45,'pale_blue')
c.text('Refined skill K*',1.81,2.68,2.18,.31,20,True)
c.text('Graph annotations + skill resources',.38,3.28,3.80,.26,14,False,PP_ALIGN.CENTER)
c.chip('Freeze for evaluation',.74,3.64,3.08,.25,'pale_teal',14)

# Reading order: top left-to-right; down from compilation; bottom right-to-left.
for j in range(3):full.arrow(XS[j]+PW+.055,YTOP+1.88,.25,.25)
for j in range(3):full.arrow(XS[j]+PW+.055,YBOT+1.88,.25,.25,'left')
full.arrow(XS[3]+PW/2-.16,5.03,.32,.99,'down')
full.text('Initial skill K⁰',XS[3]+PW/2+.26,5.65,1.90,.29,15,True)
# The graph/evidence remain available to GRAFT, separately from text consumed by the agent.
full.line(XS[1]+PW/2,4.98,XS[1]+PW/2,5.25,'navy',1.5,dash=True)
full.line(XS[1]+PW/2,5.25,XS[2]+PW/2,5.25,'navy',1.5,dash=True)
full.line(XS[2]+PW/2,5.25,XS[2]+PW/2,6.12,'navy',1.5,True,True)
full.text('Graph + evidence',XS[2]+PW/2+.18,5.93,2.0,.18,12,False)

slide.notes_slide.notes_text_frame.text=(
    'Figure 2 overview. Upper row: offline structural prior induction, steps 1–4 left to right. '
    'Lower row: online refinement, steps 5–8 right to left. '
    'The handoff is the compiled initial skill; the observed graph and evidence index remain available for GRAFT. '
    'The graph is intermediate; the fixed agent consumes text skill resources. '
    'Training-only feedback; freeze the final artifact for held-out evaluation. '
    'Avatar PNGs are the only raster objects. Modules, graphs, arrows, documents, labels and icons are native editable shapes.')


def clone_group(source, dest):
    element=copy.deepcopy(source._element)
    # Images in each new slide need destination-local relationship IDs.
    for node in element.iter():
        for attr,val in list(node.attrib.items()):
            if attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                rel=slide.part.rels[val]
                node.set(attr,dest.part.relate_to(rel.target_part if not rel.is_external else rel.target_ref,
                                                rel.reltype,is_external=rel.is_external))
    dest.shapes._spTree.insert_element_before(element,'p:extLst')
    return dest.shapes[-1]


for n in range(1,9):
    detail=prs.slides.add_slide(prs.slide_layouts[6])
    cc=Canvas(detail.shapes)
    cc.text('SkillMiner · Fig. 2 / editable component',.45,.22,16,.36,20,True)
    g=clone_group(MODULES[n],detail)
    factor=2.20
    g.left=Inches((W-PW*factor)/2);g.top=Inches(.91)
    g.width=Inches(PW*factor);g.height=Inches(PH*factor)
    detail.notes_slide.notes_text_frame.text=NOTES[n]

prs.save(PPTX)

# Structural verification of the delivered PPTX.
check=Presentation(PPTX)
assert len(check.slides)==9
for i,s in enumerate(check.slides):
    for sh in s.shapes:
        assert sh.left>=0 and sh.top>=0,(i,sh.name,'negative')
        assert sh.left+sh.width<=check.slide_width+100,(i,sh.name,'right overflow')
        assert sh.top+sh.height<=check.slide_height+100,(i,sh.name,'bottom overflow')
assert len([s for s in check.slides[0].shapes if s.shape_type==6])==8

manifest={'pptx':PPTX.name,'slides':9,'canvas_inches':[W,H],
          'layout':'offline left-to-right; online right-to-left',
          'native_editable':'all diagrams, arrows, labels, panels and document pictograms',
          'raster_assets':'customer/service avatar PNGs only',
          'source_method':'paper/overleaf_sync/sections/3_methodology.tex',
          'steps':[{'id':n,'title':TITLES[n],'notes':NOTES[n]} for n in range(1,9)]}
(OUT/'figure_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print('Created:',PPTX)
print('Verified: 9 slides, 8 editable module groups, all top-level shapes within canvas.')


def export():
    import win32com.client
    import fitz
    app=win32com.client.DispatchEx('PowerPoint.Application')
    presentation=None
    try:
        presentation=app.Presentations.Open(str(PPTX),ReadOnly=True,Untitled=False,WithWindow=False)
        allpdf=OUT/'SkillMiner_Fig2_working_pages_v1.pdf'
        presentation.SaveAs(str(allpdf),32)
        preview=OUT/'SkillMiner_Fig2_v1.png'
        presentation.Slides(1).Export(str(preview),'PNG',3200,1704)
        for n in range(1,9):
            presentation.Slides(n+1).Export(str(OUT/f'step_{n:02d}_preview.png'),'PNG',1800,959)
        with fitz.open(allpdf) as pdf:
            single=fitz.open();single.insert_pdf(pdf,from_page=0,to_page=0)
            single.save(OUT/'SkillMiner_Fig2_v1.pdf',garbage=4,deflate=True)
            single.close()
        print('Exported actual PowerPoint rendering: PNG, single-page PDF, and 9-page working PDF.')
    finally:
        if presentation is not None:presentation.Close()
        app.Quit()


if __name__=='__main__':
    export()
