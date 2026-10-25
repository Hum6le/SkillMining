"""Add a concise JTC algorithm to the report's existing unified body box."""
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '.pptx_deps'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement

OUT = Path(__file__).resolve().parent
src = OUT / 'SkillMiner_导师汇报_20261014_统一正文黑色版.pptx'
dest = OUT / 'SkillMiner_导师汇报_20261014_含JTC算法框.pptx'
prs = Presentation(src)
slide = prs.slides[8]
body = next(s for s in slide.shapes if s.name == 'Body')
body.left, body.top = Inches(.78), Inches(2.0)
body.width, body.height = Inches(11.85), Inches(5.12)
body.fill.solid(); body.fill.fore_color.rgb = RGBColor(255,255,255)
body.line.color.rgb = RGBColor(0,0,0); body.line.width = Pt(1)
tf = body.text_frame
tf.clear(); tf.word_wrap = True
tf.margin_left = tf.margin_right = Inches(.18)
tf.margin_top = tf.margin_bottom = Inches(.12)
rows = [
 ('Algorithm 1 · Joint TraceCover (JTC)', True),
 ('输入：训练动作图、边证据权重、覆盖目标与搜索预算', False),
 ('1. 初始化：从加权主干出发，补边至覆盖目标', False),
 ('2. 固定候选池：初始主干边 + 支持度达标的观察边', False),
 ('3. 在预算内尝试删边 / 换边：', False),
 ('    检查覆盖；用目标下界筛掉无望改进的候选', False),
 ('    对短名单求最优主干，同时检查根可达性', False),
 ('    接受满足约束且 J(G) 严格下降的最佳候选', False),
 ('4. 无短名单改进或预算耗尽时停止', False),
 ('输出：保留图 G、主干 B、补充分支 E_R', True),
 ('目标：J(G) = C(G) − β W_tree(G)，覆盖 Fρ(G) ≥ η', False),
]
for i,(line,bold) in enumerate(rows):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.line_spacing = 1.05; p.space_after = Pt(4)
    r = p.add_run(); r.text = line
    r.font.name = 'Microsoft YaHei'; r.font.size = Pt(22 if i==0 else 21)
    r.font.bold = bold; r.font.color.rgb = RGBColor(0,0,0)
    ea = OxmlElement('a:ea'); ea.set('typeface','Microsoft YaHei')
    r._r.get_or_add_rPr().append(ea)
slide.notes_slide.notes_text_frame.text += '\n\n算法框对应论文 Algorithm 1。初始化后固定候选池，主干边也允许删除。短名单中的候选均检查覆盖和根可达性，并以精确目标选择最佳可接受改动。主干会随着保留图变化，不是冻结树。停止只表示预算耗尽或短名单内没有改进，不表示全局最优。'
prs.save(dest)
check = Presentation(dest)
assert len(check.slides) == 16
for s in check.slides:
    texts = [x for x in s.shapes if x.has_text_frame and x.text.strip()]
    assert len(texts) == 2
    for x in texts:
        for p in x.text_frame.paragraphs:
            for r in p.runs: assert r.font.color.rgb == RGBColor(0,0,0)
print('Verified: JTC algorithm on slide 9; all slides retain one title + one black body text box.')
print(dest)
