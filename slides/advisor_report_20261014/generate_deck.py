"""Native, editable Chinese advisor report based on the current paper draft.

No figures, images, tables, or charts are embedded. Existing decks are retained.
"""
from pathlib import Path
import json
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '.pptx_deps'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image, ImageDraw, ImageFont

OUT = Path(__file__).resolve().parent
FONT = 'Microsoft YaHei'
NAVY, BLUE, GRAY, LIGHT = '203247', '2467A8', '667789', 'EEF4FA'

SLIDES = [
 dict(section='阶段汇报 · 目标投稿 WWW', title='SkillMiner', subtitle='面向 ToD agent 的可复用过程技能挖掘',
      points=[], takeaway='从历史对话发现结构，再用执行反馈完善技能。',
      notes='这次汇报沿用当前论文结构：先解释客服任务为什么需要技能，再讲离线挖掘和在线修订，最后介绍已有结果与未完成验证。当前论文仍是阶段稿，SGD、部分基线及机制消融尚未补齐。'),
 dict(section='动机', title='客服 agent 需要可维护的操作流程',
      points=[('任务有程序性', '验证身份、检查订单、执行服务操作'),
              ('规则需要更新', '政策变化会改变条件、顺序和处理方式'),
              ('技能便于维护', '复用、检查和修改独立于模型参数')],
      takeaway='我们希望把程序性知识保存为可修改的 skill。',
      notes='ToD 任务除了生成回复，还要按政策完成一系列服务操作。业务政策具有可更新性，因此知识如果只隐含在模型参数或历史对话里，维护成本就很高。外部技能给我们一个可检查、可修改的知识载体，但技能本身并不构成合规保证。'),
 dict(section='动机', title='对话不同，背后往往共享一段流程',
      points=[('共同步骤', '查账户 → 验证订单 → 检查状态'),
              ('条件分支', '退款请求与取消请求需要不同操作'),
              ('关键观察', '用户交互提供分支条件的证据')],
      takeaway='共享主干与条件分支，应当一起被保存。', reserve='待插入：Figure 1 对话与 skill 示例',
      notes='用 shipping_issue 的示意案例解释：多段对话都可能经过账户、订单和物流检查，但用户请求与订单状态不同，下一步操作就不同。这个例子强调共享结构和条件分支；它是改写案例，不是数据集给出的官方 SOP。'),
 dict(section='相关工作', title='经验复用已有基础，结构仍需显式发现',
      points=[('Skill 挖掘', 'Trace2Skill、SKILL-DISCO：归纳可复用程序'),
              ('Memory 学习', 'AWM、ExpeL、A-MEM：积累执行经验'),
              ('Skill 演化', '从已有技能出发，利用执行反馈修订')],
      takeaway='我们的入口：历史对话 → 结构先验 → 文本 skill。',
      notes='相关工作包括 ToD 程序知识、skill induction、skill evolution 和 memory。这里为了简洁把 ToD 程序背景放在动机部分。Trace2Skill 整合轨迹经验，SKILL-DISCO 使用参数化状态机子图，AWM、ExpeL、A-MEM 关注经验记忆；SkillRevise、SelfEvoSkill、SkillProx 关注已有技能的修订。部分演化方法也能生成初始技能，因此区别要落在我们从对话发现共享结构，并用这个结构连接初始化与更新，而不是笼统声称它们都依赖人工技能。'),
 dict(section='贡献', title='一个结构，连接技能初始化与后续修订',
      points=[('离线：发现结构先验', '无需人工 skill 标注，可接入固定 agent'),
              ('在线：执行反馈补充结构', '图定位问题，修订条件与操作规则'),
              ('验证：执行成功率', '用完整操作及参数正确性检验效果')],
      takeaway='结构既组织技能，也组织反馈。',
      notes='第一项贡献强调可插拔的离线算法：从历史对话构造主干和补充证据，不需要 agent 训练 rollout 或模型参数更新，但仍可能调用语言模型归纳条件和编译文本。第二项贡献是利用同一图结构组织执行反馈，并补充条件、规则和参数政策。第三项贡献是已有 ABCD 结果中的执行性能。'),
 dict(section='方法总览', title='离线挖掘与在线修订构成两阶段框架',
      points=[('离线挖掘', '历史对话 → 图结构 → 初始 skill'),
              ('在线修订', '训练执行 → 局部诊断 → 更新 skill'),
              ('最终评估', '冻结技能，在留出数据上测试')],
      takeaway='“在线”使用训练反馈，测试阶段保持技能冻结。', reserve='待插入：离线 / 在线方法总览图',
      notes='整体上先离线形成初始技能，再运行固定 agent 获得训练执行反馈，依据局部诊断修订技能。当前论文里的 online 是训练阶段的反馈修订，而不是在测试过程中读答案或实时学习测试反馈。这个边界要在汇报时说清楚。'),
 dict(section='方法 · 3.1', title='最终交付的是文本技能包',
      points=[('入口文档 T', '简洁流程、路由条件、资源读取说明'),
              ('资源集合 R', '详细动作规则、参数政策、参考案例'),
              ('中间结构 G', '组织挖掘、保留来源、定位后续更新')],
      takeaway='目标：让固定 agent 更可靠地完成未来决策。',
      notes='形式上 skill 是 K=(T,R)，不是单独的一张图。图是中间挖掘与更新状态，agent 实际读取编译后的文档。我们的学习目标是降低未来决策的失败率；失败意味着动作或必要参数有任一不正确。'),
 dict(section='方法 · 3.2', title='把动作变成节点，把交互保存在边上',
      points=[('节点与边', '节点是服务动作；边是相邻动作转移'),
              ('上下文证据', '保留两次动作之间的客户与客服交互'),
              ('分支依据', '比较同一动作的不同后继及其语境')],
      takeaway='用户话语不作节点，但参与条件归纳。',
      notes='每段对话被投影为有序动作序列，相邻动作形成有向边。边同时关联动作之间的客户与客服话语、目标动作之前的历史和参数。这样不把客户话语画成动作节点，也能利用它们解释后继动作为什么不同。重复动作可形成重试证据，虚拟根只负责组织入口。'),
 dict(section='方法 · 3.2', title='JTC 联合选择主干与补充转移',
      points=[('覆盖多样轨迹', '避免高频路径掩盖较少见的流程'),
              ('控制图的复杂度', '减少冗余边与过多候选分支'),
              ('支持清晰主干', '偏好证据充分、便于组织的动作骨架')],
      takeaway='先构造可行图，再在覆盖约束下删边或换边。',
      notes='Joint TraceCover 的目标是在覆盖和可达性约束下，最小化 C(G)-β W_tree(G)。先由证据加权主干和按覆盖增益加入的补充边得到初始可行图，再用有预算的局部搜索删边或换边，每次重新考虑最优主干。被接受的移动保持约束并使目标下降；全局近似差距的量化证明仍是论文中的待补内容，当前不能宣称已经给出近似保证。'),
 dict(section='方法 · 3.3', title='把结构编译成 agent 能读的知识',
      points=[('归纳继续条件', '联合比较后继案例，标注分支与不确定性'),
              ('生成执行规则', '写清动作要求、参数来源和缺失处理'),
              ('按需读取细节', '入口保留流程，资源保存完整规则与证据')],
      takeaway='主流程简洁，细节与例外仍可检索。',
      notes='选中的拓扑还不足以执行，因为动作名字没有说明具体条件和参数。我们先收集转移案例和参数位置统计，再比较同一源动作的所有观察后继，归纳可区分的条件、回退关系或未充分说明的分支。完整资源按需读取，示例中的具体取值不会作为固定默认值写入政策。'),
 dict(section='方法 · 3.4', title='诊断时先区分“选错”与“做错”',
      points=[('选错操作', '检查分支路由、动作选择及竞争后继'),
              ('操作细节错误', '动作正确时，检查参数政策与资源读取'),
              ('图邻域内对照', '把成功与失败案例一起用于定位')],
      takeaway='局部图提供问题位置，执行记录提供修改依据。',
      notes='我们保留同一初始技能快照下的训练执行记录。动作选错属于选择错误；只有动作正确时才用参数不匹配诊断细节错误。记录关联参考转移、竞争动作和局部资源，诊断批次可包含成功与失败。诊断报告仍是假设，需要后续证据整合，不能直接当成有效修改的证明。'),
 dict(section='方法 · 3.5', title='合并诊断证据，再修改技能',
      points=[('整合相邻报告', '合并兼容建议，检查冲突与弱证据'),
              ('修订相关知识', '更新分支条件、动作规则、参数政策'),
              ('保留成功行为', '用已成功案例检验修改是否过于宽泛')],
      takeaway='图与文档同步更新，最终版本冻结后再评估。',
      notes='不同批次可能描述同一问题，所以按图 footprint 和语义关系整合报告。成功案例可反驳过宽的修改建议，支持条件清楚的分支可以进入可见流程。合法引用和文本匹配检查能保证可应用性与记录一致性，但不构成不回退的形式化保证。'),
 dict(section='实验 · 4.1', title='成功率要求动作与参数同时正确',
      points=[('数据范围', 'ABCD：10 个流程；SGD：16 个训练域家族'),
              ('统一比较', '固定 ReAct agent、相同模型与评估协议'),
              ('指标与成本', 'SR、动作/参数准确率、调用数、token 用量')],
      takeaway='知识只用训练数据构建；测试前冻结。',
      notes='ABCD 全量约一万段对话，按十个高层流程组织。SGD 保留官方划分，合并版本化服务为域家族，只覆盖训练里出现的十六个家族。SR 是单回合所有动作及相关参数同时正确的比例。ABCD 参数按顺序精确比较，SGD 按对应动作集、API 参数和对话行为字段比较。基线包括 ReAct、Reflexion、AWM、ExpeL、A-MEM、Trace2Skill 和 SKILL-DISCO。各方法默认超参数并外接 ReAct，具体配置仍需最终审计；mining 和测试执行成本分开记录。'),
 dict(section='实验 · 4.2', title='ABCD 已出结果中，完整系统表现最好',
      points=[('成功率 50.23%', '比 Trace2Skill 高 4.89 个百分点'),
              ('Action 63.11%｜Slot 60.09%', '均高于当前已出结果的基线'),
              ('对比尚未齐全', 'ReAct、A-MEM 和 SGD 结果待补')],
      takeaway='结论限定于当前已完成的对比。', reserve='待插入：ABCD / SGD 主结果表',
      notes='主实验完整系统的 SR 为 50.23%，Trace2Skill 为 45.34%，差距为 4.89 个百分点。Trace2Skill 这版已完成。Action accuracy 为 63.11%，slot accuracy 为 60.09%。当前表中最优 SkillMiner 版本 mining calls 为 91.75k、tokens 为 7.141×10^8；比前一在线版本调用数更低但 token 更高，因此不能把版本变化概括成全面节省成本。ReAct、A-MEM 和 SGD 仍待补。主表不同版本不单独构成控制消融。'),
 dict(section='实验 · 4.3', title='A2：两阶段结合优于单独在线修订',
      points=[('shipping_issue：916 个 turns', '当前是单个 subflow 的结果'),
              ('Full 42.47%｜Online 31.99%', '两阶段结合高出 10.48 个百分点'),
              ('Offline 5.46%｜Schema 2.62%', 'JTC + AWM / Trace2Skill 仍待跑')],
      takeaway='这组结果支持该 subflow 上的组合收益。', reserve='待插入：A2 两阶段消融表',
      notes='A2 四组依次是 Schema only、Offline only、Online only 和 Full，其 SR 为 2.62%、5.46%、31.99%、42.47%。Full 比 online only 高 10.48 个百分点，slot accuracy 从 35.92% 到 45.09%，action accuracy 从 57.42% 到 59.72%。离线-only 的低 SR 提醒我们初始结构并不自动带来完整执行能力。这个观察局限于 916 个 action turns 的单个 subflow，不能直接推广到所有流程。'),
 dict(section='后续验证 · 4.3', title='接下来验证结构与更新机制各自的作用',
      points=[('A3：离线结构', 'Flat／Full graph／主干+补充／Random backbone'),
              ('A4：在线图信息', '局部图／去邻域／去图分组／同时去掉'),
              ('补齐论文证据', '接入 AWM / Trace2Skill，补 SGD 与近似分析')],
      takeaway='希望讨论：贡献表述是否清楚，实验是否形成闭环？',
      notes='A3 固定动作和参数 cards、agent 与读取上限，比较覆盖范围和布局方式，实际 token 仍可不同。A4 从同一初始 skill 和缓存训练反馈出发，用同一 updater 检查显式图邻域与分组；graph-local 也要在同一受控 replay 下重跑，不能直接搬用 A2 Full。去邻域并不删除初始 skill 中已有的图知识。目前 A3、A4 和 JTC 接其他更新方法的结果都待补，量化近似证明也待完成。希望导师重点讨论结构挖掘的贡献定位和这些验证能否闭环。'),
]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
prs.core_properties.title = 'SkillMiner：面向 ToD agent 的可复用过程技能挖掘'
prs.core_properties.subject = '导师阶段汇报；基于截至 2026-10-14 的论文草稿'
prs.core_properties.language = 'zh-CN'
PREVIEW = []

def text(slide, x, y, w, h, value, size=24, color=NAVY, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(0)
        p.line_spacing = 1.12
        run = p.add_run(); run.text = line
        run.font.name = FONT; run.font.size = Pt(size)
        run.font.bold = bold; run.font.color.rgb = RGBColor.from_string(color)
        rp = run._r.get_or_add_rPr()
        ea = OxmlElement('a:ea'); ea.set('typeface', FONT); rp.append(ea)
    PREVIEW[-1].append((x, y, w, h, value, size, color, bold))
    return shape

for idx, item in enumerate(SLIDES, 1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(255,255,255)
    PREVIEW.append([])
    text(slide, .74, .40, 11.8, .35, item['section'], 16, BLUE, True)
    if idx == 1:
        text(slide, .78, 1.70, 11.8, 1.0, item['title'], 48, NAVY, True)
        text(slide, .78, 2.88, 11.8, .8, item['subtitle'], 30)
        text(slide, .78, 4.35, 11.8, .6, item['takeaway'], 26, BLUE)
        text(slide, .78, 5.7, 11.8, .4, '导师汇报  /  2026.10.14', 20, GRAY)
    else:
        text(slide, .74, 1.03, 11.85, .75, item['title'], 30, NAVY, True)
        width = 7.7 if item.get('reserve') else 11.8
        for j, (head, detail) in enumerate(item['points']):
            y = 2.12 + j * 1.28
            text(slide, .80, y, width, .46, head, 24, BLUE, True)
            text(slide, .80, y + .48, width, .72, detail, 22, NAVY)
        if item.get('reserve'):
            text(slide, 9.15, 3.50, 3.1, 1.0, item['reserve'], 18, 'A2ADB8')
        text(slide, .78, 6.34, 11.85, .65, item['takeaway'], 21, BLUE, True)
    text(slide, .78, 7.10, 11.4, .25, 'SkillMiner  ·  阶段稿', 14, GRAY)
    text(slide, 12.05, 7.08, .6, .28, f'{idx:02d}', 14, GRAY)
    slide.notes_slide.notes_text_frame.text = item['notes']

pptx_path = OUT / 'SkillMiner_导师汇报_20261014.pptx'
prs.save(pptx_path)
(OUT / 'slide_content.json').write_text(json.dumps(SLIDES, ensure_ascii=False, indent=2), encoding='utf-8')

lines = ['# SkillMiner 导师汇报：大纲与讲稿', '',
         '16 页，建议讲述约 12–15 分钟。', '',
         '来源：paper/overleaf_sync/main.tex（截至 2026-10-14）。',
         'PPT 中所有正文与预留提示均可编辑；未嵌入图片、表格或图表。', '']
for idx, item in enumerate(SLIDES, 1):
    lines.extend([f'## {idx}. {item["title"]}', '', item['notes'], ''])
    if item.get('reserve'):
        lines.extend([item['reserve'], ''])
(OUT / '大纲与讲稿.md').write_text('\n'.join(lines), encoding='utf-8')
(OUT / 'SLIDES_STATE.json').write_text(json.dumps({
    'status':'completed', 'slide_count':len(SLIDES), 'language':'zh-CN',
    'source':'paper/overleaf_sync/main.tex', 'source_commit':'4ce5502',
    'format':'editable native PowerPoint', 'images':0, 'tables':0,
    'external_review':'unavailable; local content and layout review',
    'pdf_export':'not requested; no local Office/TeX renderer available',
}, ensure_ascii=False, indent=2), encoding='utf-8')

# A contact sheet from the same text coordinates supports local layout review.
# It is a separate preview asset and is never embedded in the deck.
scale = 80
font_file = Path('C:/Windows/Fonts/msyh.ttc')
tiles = []
for idx, objects in enumerate(PREVIEW, 1):
    canvas = Image.new('RGB', (1067, 600), 'white')
    draw = ImageDraw.Draw(canvas)
    for x,y,w,h,value,size,color,bold in objects:
        font = ImageFont.truetype(str(font_file), round(size * scale / 72))
        row, wrapped = '', []
        for segment in value.split('\n'):
            row = ''
            for ch in segment:
                if row and draw.textlength(row + ch, font=font) > w * scale:
                    wrapped.append(row); row = ch
                else: row += ch
            wrapped.append(row)
        draw.multiline_text((x*scale,y*scale), '\n'.join(wrapped), font=font,
                            fill='#'+color, spacing=round(size*scale/72*.2))
    canvas.save(OUT / f'preview_{idx:02d}.png')
    tiles.append(canvas.resize((533,300)))
sheet = Image.new('RGB',(2132,1200),'#DCE3EB')
for i,tile in enumerate(tiles): sheet.paste(tile,((i%4)*533,(i//4)*300))
sheet.save(OUT / 'overview.png')

# Validate the delivered file, not only the in-memory representation.
check = Presentation(pptx_path)
assert len(check.slides) == 16
for idx, slide in enumerate(check.slides, 1):
    assert slide.notes_slide.notes_text_frame.text.strip(), f'notes missing: {idx}'
    for shape in slide.shapes:
        assert not shape.has_table, f'table embedded: {idx}'
        assert shape.shape_type != 13, f'image embedded: {idx}'
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= check.slide_width
        assert shape.top + shape.height <= check.slide_height
print(f'Created {pptx_path}')
print('Verified: 16 slides, speaker notes, editable text, no images/tables, shapes within slide bounds.')
